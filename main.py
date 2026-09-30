import os
import sys
import argparse
import asyncio
import base64
import io
import re
import zipfile
import urllib.parse
import concurrent.futures
import threading
import json
import unicodedata
import html
import atexit
import time
from bs4 import BeautifulSoup
from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeoutError
from googleapiclient.discovery import build
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from email.header import decode_header
from datetime import datetime, timedelta

from quet_email import run_scan, list_all_messages, get_local_service, GMAIL_RETRIES, SCAN_WORKERS
from email_state import (
    EmailState, DONE, FAILED, SKIPPED, PENDING, ALL_STATUSES, gmail_link,
    norm_code, norm_so, in_master, PETRO_NOT_FOUND_MAX,
)

try:
    import pdfplumber
except Exception:
    pdfplumber = None

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

# Nhiều luồng tải cùng in log -> khóa lại để mỗi lệnh print ra trọn 1 dòng, không dính vào nhau.
_PRINT_LOCK = threading.RLock()
_builtin_print = print
def print(*args, **kwargs):
    with _PRINT_LOCK:
        _builtin_print(*args, **kwargs)

DOWNLOAD_DIR = "downloads"
TOKEN_PATH = "token.json"
CREDENTIALS_PATH = "credentials.json"
LABEL_NAME = "xd"
SOURCE = "inbox"          # inbox | all | label
EXTRA_QUERY = ""          # điều kiện Gmail bổ sung, vd: has:attachment
SINCE_TS = None           # chỉ quét mail từ mốc thời gian này (epoch giây)
UNTIL_TS = None           # chỉ quét mail TRƯỚC mốc này (epoch giây) = hết ngày --until
STATE_PATH = os.path.join("state", "processed_default.db")
MAX_ATTEMPTS = 3
PETRO_OUTPUT = "hoa_don_petrolimex.xlsx"   # file Excel Petrolimex
PETRO_KEYWORDS = ["petrolimex"]            # từ khóa nhận diện mail Petro (người gửi/tiêu đề/nội dung)
PETRO_LABEL = "petro"                      # label Gmail chứa mail Petro ("" = không dùng)
PETRO_DOWNLOAD_DIR = "HoaDon_Petrolimex"   # thư mục lưu PDF Petro (mặc định nằm trong thư mục tải của profile)
PETRO_MASTER = ""                          # file Excel TỔNG các hóa đơn Petro đã tải (dùng lọc trùng); "" = tắt
PETRO_MASTER_DIR = r"D:\Làm việc\hđ đang xử lý"      # thư mục chứa file tổng Petro
PETRO_MASTER_FILENAME = "hoa_don_petrolimex_tong_hop_down.xlsx"   # tên file tổng Petro
PETRO_LIMIT = 0                            # số hóa đơn Petro tải tối đa mỗi lần (0 = tất cả)
PETRO_CAPTCHA_TIMEOUT = 90                 # số giây chờ bạn nhập CAPTCHA cho mỗi mã
PETRO_START_DELAY = 10                     # số giây đếm ngược trước khi tự mở trình duyệt tải Petro
PETRO_LOOKUP_URL = "https://hoadon.petrolimex.com.vn/"
PETRO_AUTO_CAPTCHA = True                  # tự giải CAPTCHA bằng ddddocr (tắt: --petro-manual-captcha)
PETRO_CAPTCHA_ATTEMPTS = 5                 # số lần tự giải CAPTCHA mỗi mã trước khi chuyển nhập tay
PETRO_HEADLESS = False                     # True = chạy ẩn, chỉ tự giải, không nhập tay
PETRO_XML = False                          # tải kèm file XML hóa đơn Petro
KEEP_XML = False                           # lưu thêm file XML hóa đơn (đính kèm email / nằm trong file ZIP) - bật bằng --keep-xml
XML_SUBDIR = "XML"                         # thư mục con (trong thư mục tải của profile) chứa file XML
ARGS_PROFILE = ""                          # tên profile đang chạy (chỉ dùng để in gợi ý lệnh)
MAX_CONCURRENT = 3
SCOPES = ["https://www.googleapis.com/auth/gmail.readonly"]
LOG_DIR = "logs"                           # thư mục lưu log mỗi lần chạy
LOG_PATH = ""                              # file log của lần chạy hiện tại (tự sinh)
ERROR_DIR = "outputs"                      # thư mục lưu file Excel lỗi/bỏ qua/chưa xử lý (tự sinh theo profile)

os.makedirs(DOWNLOAD_DIR, exist_ok=True)
_FILENAME_LOCK = threading.Lock()
_RESERVED_PATHS = set()

# ================== COMMON ==================
def ensure_dir(path):
    if not os.path.exists(path):
        os.makedirs(path)

# ================== LOG CHẠY (ghi song song ra màn hình và file) ==================
class _LogFile:
    """Ghi log ra file, thêm mốc giờ [HH:MM:SS] đầu mỗi dòng. Dùng chung cho stdout và stderr."""
    def __init__(self, fh):
        self._fh = fh
        self._new_line = True
        self._lock = threading.Lock()

    def write(self, text):
        if not text or text.startswith("\r"):   # bỏ qua dòng đếm ngược (\r) để log gọn
            return
        with self._lock:
            for part in text.splitlines(keepends=True):
                if self._new_line:
                    self._fh.write(datetime.now().strftime("[%H:%M:%S] "))
                self._fh.write(part)
                self._new_line = part.endswith("\n")

    def flush(self):
        with self._lock:
            self._fh.flush()

class _Tee:
    def __init__(self, stream, logfile):
        self._stream = stream
        self._logfile = logfile

    def write(self, text):
        n = self._stream.write(text)
        try:
            self._logfile.write(text)
        except Exception:
            pass            # lỗi ghi log không được làm hỏng chương trình
        return n

    def flush(self):
        self._stream.flush()
        try:
            self._logfile.flush()
        except Exception:
            pass

    def __getattr__(self, name):
        return getattr(self._stream, name)

def setup_logging(log_dir, profile):
    """Bắt đầu ghi toàn bộ những gì in ra màn hình (kể cả lỗi) vào file log riêng của lần chạy."""
    global LOG_PATH
    if LOG_PATH:
        return
    try:
        os.makedirs(log_dir, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = os.path.join(log_dir, f"run_{profile or 'default'}_{stamp}.log")
        fh = open(path, "a", encoding="utf-8", buffering=1)
    except Exception as e:
        print(f"[WARN] Không tạo được file log trong '{log_dir}' | ERR={repr(e)}")
        return
    LOG_PATH = path
    logfile = _LogFile(fh)
    sys.stdout = _Tee(sys.stdout, logfile)
    sys.stderr = _Tee(sys.stderr, logfile)
    atexit.register(fh.close)
    print(f"[LOG] Ghi log tại: {LOG_PATH}")
    print(f"[LOG] Bắt đầu: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Lệnh: {' '.join(sys.argv)}")

def safe_profile_name(profile):
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", profile.strip())

def _resolve_path(path):
    if not path:
        return path
    if os.path.isabs(path):
        return path
    return os.path.join(os.getcwd(), path)

def load_app_config():
    config_path = os.path.join(os.getcwd(), "app_config.json")
    if not os.path.exists(config_path):
        return {}
    try:
        with open(config_path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        print(f"[WARN] Khong doc duoc app_config.json | ERR={repr(e)}")
        return {}

def find_profile_config(profile):
    if not profile:
        return None
    for item in load_app_config().get("profiles", []):
        if str(item.get("id", "")).lower() == profile.lower():
            return item
    return None

def parse_args():
    parser = argparse.ArgumentParser(description="Tải hóa đơn từ Gmail (Inbox hoặc label), có sổ theo dõi email đã xử lý.")
    parser.add_argument("--profile", help="Tên profile email, ví dụ: s7, lh, thp.")
    parser.add_argument("--download-dir", help="Thư mục lưu file tải về.")
    parser.add_argument("--token", help="Đường dẫn token riêng cho email/profile.")
    parser.add_argument("--credentials", default="credentials.json", help="Đường dẫn credentials.json.")
    parser.add_argument("--label", default="xd", help="Tên label Gmail (dùng cho --source label và --seed-label).")
    parser.add_argument("--keep-xml", action="store_true", help="Lưu thêm file XML hóa đơn (đính kèm email hoặc nằm trong file ZIP) vào thư mục con XML. Không ảnh hưởng việc tải PDF.")

    g = parser.add_argument_group("Nguồn email")
    g.add_argument("--source", choices=["inbox", "all", "label"],
                   help="inbox (mặc định) = hộp thư đến | all = mọi thư (kể cả đã archive, trừ Sent/Nháp) | label = theo label.")
    g.add_argument("--query", help="Điều kiện Gmail bổ sung, ví dụ: \"has:attachment\" hoặc \"from:(meinvoice.vn)\".")
    g.add_argument("--since", help="Chỉ quét mail từ ngày này (YYYY-MM-DD), gồm cả ngày này. Nên dùng cho lần chạy đầu tiên.")
    g.add_argument("--until", help="Chỉ quét mail đến hết ngày này (YYYY-MM-DD), gồm cả ngày này. Ví dụ: --since 2026-09-19 --until 2026-09-21")

    g = parser.add_argument_group("Sổ theo dõi (email đã quét/xử lý)")
    g.add_argument("--state", help="Đường dẫn file sổ theo dõi (.db). Mặc định: state/processed_<profile>.db")
    g.add_argument("--max-attempts", type=int, default=3, help="Số lần thử tối đa cho email lỗi (mặc định 3).")
    g.add_argument("--retry-failed", action="store_true", help="Thử lại cả email lỗi đã hết lượt thử.")
    g.add_argument("--reprocess-skipped", action="store_true", help="Quét lại cả email đã bị SKIPPED (vd. sau khi thêm supplier mới).")
    g.add_argument("--force", action="store_true", help="Bỏ qua sổ theo dõi, quét & tải lại tất cả (có thể tạo file trùng).")
    g.add_argument("--seed-label", metavar="LABEL", help="Đánh dấu toàn bộ mail đang có label này là DONE (dùng 1 lần khi chuyển từ label sang Inbox).")

    g = parser.add_argument_group("Petrolimex (xử lý sau cùng, xuất Excel Số hóa đơn / Mã tra cứu)")
    g.add_argument("--petro-output", help="File Excel Petro. Mặc định: outputs/<profile>/hoa_don_petrolimex.xlsx")
    g.add_argument("--petro-keyword", nargs="+", metavar="KW", help="Từ khóa nhận diện mail Petro trong người gửi/tiêu đề/nội dung (mặc định: petrolimex).")
    g.add_argument("--petro-label", help="Label Gmail chứa mail Petro (mặc định: petro). Truyền \"\" để tắt.")
    g.add_argument("--export-petro", action="store_true", help="Xuất lại file Excel Petro từ sổ theo dõi (không quét Gmail).")
    g.add_argument("--no-petro-download", action="store_true", help="Không tải PDF Petro ở cuối lần chạy (bỏ qua STEP 6).")
    g.add_argument("--download-petro", action="store_true", help="CHỈ tải PDF Petro còn thiếu (không quét Gmail). Tự giải CAPTCHA, không được thì chờ nhập tay.")
    g.add_argument("--petro-download-dir", help="Thư mục lưu PDF Petro. Mặc định: <thư mục tải của profile>/HoaDon_Petrolimex")
    g.add_argument("--petro-master", help="File Excel TỔNG các hóa đơn Petro đã tải: hóa đơn có trong file này sẽ KHÔNG tải lại, hóa đơn mới tải xong được thêm vào. Mặc định: D:\\Làm việc\\hđ đang xử lý\\hoa_don_petrolimex_tong_hop_down.xlsx. Truyền \"\" để tắt.")
    g.add_argument("--petro-limit", type=int, default=0, help="Tải tối đa N hóa đơn Petro mỗi lần (0 = tất cả).")
    g.add_argument("--petro-captcha-timeout", type=int, default=90, help="Số giây chờ NHẬP TAY CAPTCHA cho mỗi mã (mặc định 90).")
    g.add_argument("--petro-captcha-attempts", type=int, default=5, help="Số lần TỰ GIẢI CAPTCHA mỗi mã trước khi chuyển nhập tay (mặc định 5).")
    g.add_argument("--petro-manual-captcha", action="store_true", help="Tắt tự giải CAPTCHA, luôn nhập tay như cách cũ.")
    g.add_argument("--petro-headless", action="store_true", help="Chạy ẩn trình duyệt, chỉ tự giải CAPTCHA (không nhập tay).")
    g.add_argument("--petro-xml", action="store_true", help="Tải kèm file XML hóa đơn Petro.")
    g.add_argument("--mark-petro-downloaded", nargs="*", metavar="KEY", help="Đánh dấu hóa đơn Petro là ĐÃ TẢI PDF (theo mã tra cứu hoặc ID email). Không truyền gì = tất cả hóa đơn đang chờ tải.")

    g = parser.add_argument_group("Kiểm tra thủ công (không cần quét Gmail)")
    g.add_argument("--report", action="store_true", help="In báo cáo tổng hợp và xuất lại file CSV.")
    g.add_argument("--status", type=str.lower, choices=["done", "failed", "skipped", "pending"], help="Dùng với --report: chỉ liệt kê trạng thái này.")
    g.add_argument("--limit", type=int, default=50, help="Số dòng tối đa khi liệt kê (0 = tất cả). Mặc định 50.")
    g.add_argument("--mark-done", nargs="+", metavar="ID", help="Đánh dấu các email này là DONE (vd. bạn đã tự tải tay).")
    g.add_argument("--mark-pending", nargs="+", metavar="ID", help="Đánh dấu các email này để xử lý lại ở lần chạy sau.")
    g.add_argument("--verify", action="store_true", help="Kiểm tra email DONE có file PDF còn trong thư mục tải không.")
    g.add_argument("--fix", action="store_true", help="Dùng với --verify: tự đánh dấu PENDING các email bị thiếu file.")
    g = parser.add_argument_group("Log & file lỗi")
    g.add_argument("--log-dir", help="Thư mục lưu log chạy. Mặc định: logs/ (mỗi lần chạy 1 file run_<profile>_<ngày giờ>.log).")
    g.add_argument("--error-dir", help="Thư mục lưu file Excel các mục lỗi/bỏ qua/chưa xử lý. Mặc định: outputs/<profile>/")
    return parser.parse_args()

def configure_runtime(args):
    global DOWNLOAD_DIR, TOKEN_PATH, CREDENTIALS_PATH, LABEL_NAME
    global SOURCE, EXTRA_QUERY, SINCE_TS, UNTIL_TS, STATE_PATH, MAX_ATTEMPTS
    global PETRO_OUTPUT, PETRO_KEYWORDS, PETRO_LABEL
    global PETRO_DOWNLOAD_DIR, PETRO_LIMIT, PETRO_CAPTCHA_TIMEOUT, ARGS_PROFILE, PETRO_MASTER
    global PETRO_AUTO_CAPTCHA, PETRO_CAPTCHA_ATTEMPTS, PETRO_HEADLESS, PETRO_XML
    global LOG_DIR, ERROR_DIR, KEEP_XML

    profile = safe_profile_name(args.profile) if args.profile else ""
    profile_config = find_profile_config(profile)
    pc = profile_config or {}

    # Log chạy: bật sớm nhất có thể để ghi lại cả các cảnh báo khi đọc cấu hình
    LOG_DIR = _resolve_path(args.log_dir or pc.get("log_dir") or "logs")
    setup_logging(LOG_DIR, profile)
    DOWNLOAD_DIR = args.download_dir or (
        profile_config.get("download_dir") if profile_config else (os.path.join("downloads", profile) if profile else "downloads")
    )
    TOKEN_PATH = args.token or (
        profile_config.get("token") if profile_config else (os.path.join("tokens", f"token_{profile}.json") if profile else "token.json")
    )
    CREDENTIALS_PATH = args.credentials if args.credentials != "credentials.json" else (
        profile_config.get("credentials") if profile_config else args.credentials
    )
    LABEL_NAME = args.label if args.label != "xd" else (
        profile_config.get("label") if profile_config else args.label
    )

    # Nguồn quét: tham số dòng lệnh > app_config.json (profile) > mặc định inbox
    SOURCE = args.source or pc.get("source") or "inbox"
    EXTRA_QUERY = args.query if args.query is not None else (pc.get("query") or "")
    since = args.since or pc.get("since")
    SINCE_TS = None
    if since:
        try:
            SINCE_TS = int(datetime.strptime(str(since), "%Y-%m-%d").timestamp())
        except ValueError:
            print(f"[WARN] --since phải có dạng YYYY-MM-DD (nhận được: {since}). Bỏ qua.")

    # Ngày kết thúc TÍNH CẢ NGÀY ĐÓ: Gmail `before:` là mốc loại trừ nên lấy 00:00 của ngày kế tiếp
    until = args.until or pc.get("until")
    UNTIL_TS = None
    if until:
        try:
            UNTIL_TS = int((datetime.strptime(str(until), "%Y-%m-%d") + timedelta(days=1)).timestamp())
        except ValueError:
            print(f"[WARN] --until phải có dạng YYYY-MM-DD (nhận được: {until}). Bỏ qua.")
    if SINCE_TS and UNTIL_TS and UNTIL_TS <= SINCE_TS:
        print(f"[ERROR] Ngày kết thúc (--until {until}) phải cùng ngày hoặc sau ngày bắt đầu (--since {since}).")
        sys.exit(1)
    MAX_ATTEMPTS = max(1, args.max_attempts)
    STATE_PATH = args.state or pc.get("state") or os.path.join("state", f"processed_{profile or 'default'}.db")

    # Petrolimex: tham số dòng lệnh > app_config.json (profile) > mặc định
    kw = args.petro_keyword or pc.get("petro_keywords") or ["petrolimex"]
    PETRO_KEYWORDS = [kw] if isinstance(kw, str) else list(kw)
    PETRO_LABEL = args.petro_label if args.petro_label is not None else pc.get("petro_label", "petro")
    PETRO_OUTPUT = args.petro_output or pc.get("petro_output") or (
        os.path.join("outputs", profile, "hoa_don_petrolimex.xlsx") if profile else "hoa_don_petrolimex.xlsx"
    )

    DOWNLOAD_DIR = _resolve_path(DOWNLOAD_DIR)
    TOKEN_PATH = _resolve_path(TOKEN_PATH)
    CREDENTIALS_PATH = _resolve_path(CREDENTIALS_PATH)
    STATE_PATH = _resolve_path(STATE_PATH)
    PETRO_OUTPUT = _resolve_path(PETRO_OUTPUT)
    ERROR_DIR = _resolve_path(args.error_dir or pc.get("error_dir") or (os.path.join("outputs", profile) if profile else "outputs"))
    _master_default = os.path.join(PETRO_MASTER_DIR, PETRO_MASTER_FILENAME)
    _pm = args.petro_master if args.petro_master is not None else pc.get("petro_master", _master_default)
    PETRO_MASTER = _resolve_path(_pm) if _pm else ""
    ARGS_PROFILE = profile
    PETRO_LIMIT = max(0, args.petro_limit)
    PETRO_CAPTCHA_TIMEOUT = max(10, args.petro_captcha_timeout)
    PETRO_AUTO_CAPTCHA = not (args.petro_manual_captcha or pc.get("petro_manual_captcha", False))
    PETRO_CAPTCHA_ATTEMPTS = max(1, args.petro_captcha_attempts)
    PETRO_HEADLESS = bool(args.petro_headless or pc.get("petro_headless", False))
    PETRO_XML = bool(args.petro_xml or pc.get("petro_xml", False))
    KEEP_XML = bool(args.keep_xml or pc.get("keep_xml", False))
    PETRO_DOWNLOAD_DIR = _resolve_path(
        args.petro_download_dir or pc.get("petro_download_dir") or os.path.join(DOWNLOAD_DIR, "HoaDon_Petrolimex")
    )

    ensure_dir(DOWNLOAD_DIR)
    ensure_dir(os.path.dirname(STATE_PATH))
    token_dir = os.path.dirname(TOKEN_PATH)
    if token_dir:
        ensure_dir(token_dir)

def unzip_and_cleanup(file_path, names_out=None, email_id=""):
    try:
        with open(file_path, 'rb') as f:
            if f.read(2) != b'PK':
                return None
    except Exception as e:
        print(f"[WARN] Không đọc được file để kiểm tra ZIP: {file_path} | ERR={e}")
        return None

    extracted = 0
    zip_path = file_path + ".zip"
    try:
        extract_dir = os.path.dirname(file_path) or DOWNLOAD_DIR
        os.rename(file_path, zip_path)
        with zipfile.ZipFile(zip_path, 'r') as zip_ref:
            for member in zip_ref.namelist():
                if member.lower().endswith('.pdf'):
                    out_name = get_unique_filename(extract_dir, os.path.basename(member))
                    out_path = os.path.join(extract_dir, out_name)
                    with zip_ref.open(member) as src, open(out_path, 'wb') as dst:
                        dst.write(src.read())
                    if names_out is not None:
                        names_out.append(out_name)
                    extracted += 1
            save_xml_from_zip(zip_ref, os.path.basename(file_path), email_id)   # --keep-xml: giữ lại XML trong ZIP (tự bắt lỗi)
        if extracted:
            os.remove(zip_path)
        else:
            os.rename(zip_path, file_path)
            print(f"[WARN] ZIP không có file PDF: {file_path}")
        return extracted
    except Exception as e:
        print(f"[WARN] Lỗi giải nén ZIP: {file_path} | ERR={e}")
        try:
            if os.path.exists(zip_path) and not os.path.exists(file_path):
                os.rename(zip_path, file_path)
        except Exception as restore_err:
            print(f"[WARN] Không khôi phục được file ZIP gốc: {file_path} | ERR={restore_err}")
        return 0

def get_unique_filename(directory, filename):
    with _FILENAME_LOCK:
        base, ext = os.path.splitext(filename)
        counter = 1
        new_filename = filename
        full_path = os.path.abspath(os.path.join(directory, new_filename))
        while os.path.exists(full_path) or full_path in _RESERVED_PATHS:
            new_filename = f"{base}_{counter}{ext}"
            full_path = os.path.abspath(os.path.join(directory, new_filename))
            counter += 1
        _RESERVED_PATHS.add(full_path)
        return new_filename

# ================== LƯU FILE XML HÓA ĐƠN (--keep-xml) ==================
# Chỉ LƯU THÊM file XML vào <thư mục tải>/XML, không thay đổi cách tải / đếm / ghi sổ PDF.
# Mọi lỗi khi lưu XML chỉ in cảnh báo, không làm email bị tính là lỗi.
_XML_LOCK = threading.RLock()      # RLock: _xml_note có thể được gọi khi đang giữ khóa
_XML_SAVED = []
_XML_STATUS = {}          # email_id -> {"kind", "files", "errors", "pdf"} để báo email nào có / không có / lỗi XML

def _xml_note(email_id, file=None, error=None, kind=None, pdf=None):
    if not email_id:
        return
    with _XML_LOCK:
        st = _XML_STATUS.setdefault(email_id, {"kind": "", "files": [], "errors": [], "pdf": []})
        if file and file not in st["files"]:
            st["files"].append(file)
        if error:
            st["errors"].append(str(error)[:300])
        if kind and not st["kind"]:
            st["kind"] = kind
        if pdf:
            st["pdf"].extend(pdf)

def _xml_after_email(email_id, kind, pdf_names=None, subject=""):
    """Gọi sau khi đã xét hết file đính kèm của 1 email: in 1 dòng nếu email không có XML."""
    if not KEEP_XML or not email_id:
        return
    _xml_note(email_id, kind=kind, pdf=pdf_names)
    st = _XML_STATUS.get(email_id) or {}
    if not st.get("files") and not st.get("errors"):
        extra = f" | PDF={', '.join(pdf_names)}" if pdf_names else ""
        print(f"[KEEP-XML] Không có XML đính kèm | EMAIL={email_id} | LOAI={kind}{extra}")

def _xml_dir():
    return os.path.join(DOWNLOAD_DIR, XML_SUBDIR)

def _is_xml_name(name):
    return str(name or "").lower().endswith(".xml")

def save_invoice_xml(data, name, source, email_id=""):
    """Lưu 1 file XML. Trùng tên và trùng nội dung thì không lưu lại. Trả về tên file đã lưu ("" nếu bỏ qua/lỗi)."""
    if not KEEP_XML or not data:
        return ""
    try:
        if not data[:600].lstrip(b"\xef\xbb\xbf").lstrip().startswith(b"<"):
            print(f"[WARN] [KEEP-XML] Bỏ qua {name}: nội dung không phải XML")
            return ""
        folder = _xml_dir()
        os.makedirs(folder, exist_ok=True)
        base = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", os.path.basename(str(name or ""))).strip(" .")
        if not base:
            base = f"{email_id or 'hoa_don'}.xml"
        if not base.lower().endswith(".xml"):
            base += ".xml"
        with _XML_LOCK:                                   # các luồng tải song song không ghi đè / lưu trùng nhau
            existing = os.path.join(folder, base)
            if os.path.exists(existing):
                try:
                    with open(existing, "rb") as f:
                        if f.read() == data:
                            print(f"[KEEP-XML] Đã có sẵn {XML_SUBDIR}/{base}, không lưu lại")
                            _xml_note(email_id, file=base)
                            return base
                except Exception:
                    pass
            out_name = get_unique_filename(folder, base)
            with open(os.path.join(folder, out_name), "wb") as f:
                f.write(data)
            _XML_SAVED.append(out_name)
        _xml_note(email_id, file=out_name)
        print(f"[KEEP-XML] Đã lưu {XML_SUBDIR}/{out_name} | NGUON={source}" + (f" | EMAIL={email_id}" if email_id else ""))
        return out_name
    except Exception as e:
        print(f"[WARN] [KEEP-XML] Không lưu được XML {name} (PDF không bị ảnh hưởng) | ERR={repr(e)}")
        _xml_note(email_id, error=f"Không lưu được {name}: {repr(e)}")
        return ""

def save_xml_from_zip(zip_ref, source, email_id=""):
    """Lưu các file XML nằm trong 1 file ZIP đang mở. Trả về số file đã lưu."""
    if not KEEP_XML:
        return 0
    saved = 0
    try:
        for member in zip_ref.namelist():
            if _is_xml_name(member) and not member.endswith("/"):
                if save_invoice_xml(zip_ref.read(member), os.path.basename(member), f"ZIP:{source}", email_id):
                    saved += 1
    except Exception as e:
        print(f"[WARN] [KEEP-XML] Không đọc được XML trong ZIP {source} | ERR={repr(e)}")
        _xml_note(email_id, error=f"Không đọc được XML trong ZIP {source}: {repr(e)}")
    return saved

def _attachment_bytes(service, msg_id, part):
    body = part.get("body") or {}
    if body.get("attachmentId"):
        att = service.users().messages().attachments().get(
            userId="me", messageId=msg_id, id=body["attachmentId"]).execute(num_retries=GMAIL_RETRIES)
        return base64.urlsafe_b64decode(att["data"])
    if body.get("data"):
        return base64.urlsafe_b64decode(body["data"])
    return b""

def save_xml_attachments(service, msg_id, payload):
    """Lưu file .xml đính kèm (và XML nằm trong file .zip đính kèm) của 1 email. Trả về số file đã lưu."""
    if not KEEP_XML:
        return 0
    saved = 0
    try:
        parts = []
        def walk(p):
            if p.get("parts"):
                for c in p["parts"]:
                    walk(c)
            else:
                parts.append(p)
        walk(payload)
        for part in parts:
            fname = part.get("filename") or ""
            mime = (part.get("mimeType") or "").lower()
            if not fname or is_pdf_part(part):
                continue                                  # chỉ xét file đính kèm có tên, bỏ qua PDF (đã tải ở trên)
            is_xml = _is_xml_name(fname) or mime in ("application/xml", "text/xml")
            is_zip = fname.lower().endswith(".zip") or mime in ("application/zip", "application/x-zip-compressed")
            if not (is_xml or is_zip):
                continue
            try:
                data = _attachment_bytes(service, msg_id, part)
                if is_xml:
                    if save_invoice_xml(data, fname, "GMAIL_XML", msg_id):
                        saved += 1
                elif data[:2] == b"PK":
                    with zipfile.ZipFile(io.BytesIO(data)) as z:
                        saved += save_xml_from_zip(z, fname, msg_id)
            except Exception as e:
                print(f"[WARN] [KEEP-XML] Không lấy được file đính kèm {fname} | EMAIL={msg_id} | ERR={repr(e)}")
                _xml_note(msg_id, error=f"Không lấy được file đính kèm {fname}: {repr(e)}")
    except Exception as e:
        print(f"[WARN] [KEEP-XML] Lỗi khi đọc file đính kèm XML | EMAIL={msg_id} | ERR={repr(e)}")
        _xml_note(msg_id, error=f"Lỗi khi đọc file đính kèm: {repr(e)}")
    return saved

def _email_payload(service, msg_id, parts=None):
    """parts = danh sách file đính kèm đã lấy sẵn khi quét (quet_email.attachment_parts) -> dùng luôn, không gọi Gmail.
    parts = None (không có sẵn) -> đọc email như cũ bằng messages.get(format="full")."""
    if parts is not None:
        return {"parts": parts}
    return service.users().messages().get(userId="me", id=msg_id, format="full").execute(num_retries=GMAIL_RETRIES)["payload"]

def _check_xml_one(msg_id, creds, kind, parts=None):
    try:
        service = get_local_service(creds)      # dùng lại service của luồng, không build() lại mỗi email
        save_xml_attachments(service, msg_id, _email_payload(service, msg_id, parts))
    except Exception as e:
        print(f"[WARN] [KEEP-XML] Không đọc được email để lấy XML | EMAIL={msg_id} | ERR={repr(e)}")
        _xml_note(msg_id, error=f"Không đọc được email: {repr(e)}")
    _xml_after_email(msg_id, kind)
    st = _XML_STATUS.get(msg_id) or {}
    if st.get("files") and kind == "Không có PDF":
        print(f"[KEEP-XML] Email không có PDF nhưng có XML | EMAIL={msg_id} | XML={', '.join(st['files'])}")

def check_xml_for_emails(items, creds, parts_by_id=None):
    """--keep-xml: lấy XML đính kèm của các email KHÔNG có PDF đính kèm (PDF tải qua link, hoặc email chỉ có XML/ZIP).
    parts_by_id: {email_id: danh sách file đính kèm lấy lúc quét} -> không phải đọc lại email."""
    if not KEEP_XML or not items:
        return
    parts_by_id = parts_by_id or {}
    print(f"[KEEP-XML] Kiểm tra XML đính kèm của {len(items)} email không có PDF đính kèm (tải qua link / không có PDF)...")
    with concurrent.futures.ThreadPoolExecutor(max_workers=8) as executor:
        list(executor.map(lambda it: _check_xml_one(it[0], creds, it[1], parts_by_id.get(it[0])), items))

def xml_report_rows(meta_by_id):
    """Danh sách email KHÔNG có XML / LỖI XML của lần chạy này (cho log tổng kết và sheet 'Thiếu XML')."""
    rows = []
    with _XML_LOCK:
        items = list(_XML_STATUS.items())
    for eid, st in items:
        if st["files"] and not st["errors"]:
            continue
        meta = meta_by_id.get(eid, {}) if meta_by_id else {}
        kind = st.get("kind") or ""
        if st["errors"]:
            status, hint = "LỖI XML", "Lưu XML bị lỗi -> chạy lại, hoặc mở email tải XML tay"
        elif kind.startswith("PDF qua link"):
            status, hint = "KHÔNG CÓ XML", "Email không kèm XML và không lấy được XML qua link -> mở email tải XML tay nếu cần"
        elif kind == "Không có PDF":
            status, hint = "KHÔNG CÓ XML", "Email không có PDF lẫn XML đính kèm"
        else:
            status, hint = "KHÔNG CÓ XML", "Email có PDF nhưng không gửi kèm XML -> mở email kiểm tra nếu cần"
        rows.append({
            "status": status, "kind": kind, "supplier": meta.get("supplier") or "",
            "date": meta.get("mail_date") or "", "subject": meta.get("subject") or "",
            "pdf": ", ".join(st.get("pdf") or []), "error": " | ".join(st["errors"])[:500],
            "saved": ", ".join(st["files"]), "hint": hint, "email_id": eid,
        })
    order = {"LỖI XML": 0}
    kind_order = {"PDF đính kèm": 0, "Không có PDF": 2}
    rows.sort(key=lambda r: (order.get(r["status"], 1), kind_order.get(r["kind"], 1), r["date"]))
    return rows

# ---- Ghi nhận file đã lưu theo từng email (để đưa vào sổ theo dõi / CSV) ----
_EMAIL_FILES = {}
_EMAIL_FILES_LOCK = threading.Lock()

def _track_files(email_id, names):
    with _EMAIL_FILES_LOCK:
        _EMAIL_FILES.setdefault(email_id, []).extend(names)

def pop_tracked_files(email_id):
    with _EMAIL_FILES_LOCK:
        return _EMAIL_FILES.pop(email_id, [])

# Giao diện chia PDF (đã tải) vào các thư mục con này sau khi ráp XML -> main.py tìm PDF cả trong đó
# (Trung_lap: PDF trùng ký hiệu + số với hóa đơn đã có XML, vd. cùng 1 hóa đơn gửi trong 2 email)
PDF_SUBDIRS = ("Cac_hang_khac", "Khong_co_XML", "Trung_lap")

def _find_pdf(name):
    """Đường dẫn file đã tải: ở thư mục tải hoặc 1 trong các thư mục con đã chia. Không thấy -> ""."""
    for sub in ("",) + PDF_SUBDIRS:
        p = os.path.join(DOWNLOAD_DIR, sub, name) if sub else os.path.join(DOWNLOAD_DIR, name)
        if os.path.exists(p):
            return p
    return ""

# ---- Không lưu trùng PDF: 2 email gửi cùng 1 file giống hệt nhau -> giữ file đã có, bỏ bản mới (không tạo "_1") ----
_DUP_LOCK = threading.Lock()
_DUP_SIZE_INDEX = None      # cỡ file -> {tên file PDF trong thư mục tải}
_DUP_HASH_CACHE = {}        # (tên, cỡ, mtime) -> sha1

def _pdf_sha1(path):
    import hashlib
    h = hashlib.sha1()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()

def _dedupe_pdf(path):
    """File PDF vừa lưu giống hệt (từng byte) 1 PDF đã có trong thư mục tải -> xóa bản mới, trả về tên file đã có.
    Không trùng -> ghi nhận file, trả về None. Mọi lỗi -> giữ nguyên file (trả về None)."""
    global _DUP_SIZE_INDEX
    try:
        folder = os.path.abspath(os.path.dirname(path) or DOWNLOAD_DIR)
        if folder != os.path.abspath(DOWNLOAD_DIR) or not path.lower().endswith(".pdf"):
            return None
        name = os.path.basename(path)
        size = os.path.getsize(path)
        with _DUP_LOCK:
            if _DUP_SIZE_INDEX is None:
                _DUP_SIZE_INDEX = {}
                for sub in ("",) + PDF_SUBDIRS:          # gồm cả PDF đã chia vào thư mục con
                    d = os.path.join(folder, sub) if sub else folder
                    if not os.path.isdir(d):
                        continue
                    for f in os.listdir(d):
                        fp = os.path.join(d, f)
                        if f.lower().endswith(".pdf") and os.path.isfile(fp):
                            _DUP_SIZE_INDEX.setdefault(os.path.getsize(fp), set()).add(os.path.join(sub, f) if sub else f)
            same_size = _DUP_SIZE_INDEX.setdefault(size, set())
            cands = sorted(n for n in same_size if n != name)
            if cands:
                mine = _pdf_sha1(path)
                for n in cands:
                    fp = os.path.join(folder, n)
                    if not os.path.isfile(fp):
                        same_size.discard(n)
                        continue
                    st = os.stat(fp)
                    key = (n, st.st_size, st.st_mtime)
                    if key not in _DUP_HASH_CACHE:
                        _DUP_HASH_CACHE[key] = _pdf_sha1(fp)
                    if _DUP_HASH_CACHE[key] == mine:
                        os.remove(path)
                        same_size.discard(name)
                        return os.path.basename(n)
            same_size.add(name)
    except Exception as e:
        print(f"[WARN] Không kiểm tra được PDF trùng: {path} | ERR={repr(e)}")
    return None

def _finalize_download(path, filename, supplier, email_id):
    """Giải nén nếu là ZIP, ghi nhận tên file, in log SUCCESS. Trả về số file đã lưu (0 nếu thất bại)."""
    names = []
    extracted = unzip_and_cleanup(path, names_out=names, email_id=email_id)
    saved_count = extracted if extracted is not None else 1
    if saved_count:
        kept = []
        for n in (names if extracted else [filename]):
            dup = _dedupe_pdf(os.path.join(os.path.dirname(path) or DOWNLOAD_DIR, n))
            if dup:
                print(f"[SKIP-DUP] {supplier.upper()} | EMAIL={email_id} | {n} giống hệt file đã có {dup} -> không lưu thêm bản trùng")
            kept.append(dup or n)
        _track_files(email_id, kept)
        print(f"[SUCCESS] {supplier.upper()} | EMAIL={email_id} | FILE={filename if extracted else kept[0]} | SAVED={saved_count}")
    return saved_count

# ================== PDF DETECT ==================
def is_pdf_part(part):
    mime = part.get("mimeType", "")
    filename = part.get("filename", "").lower()
    return mime == "application/pdf" or filename.endswith(".pdf")

# ================== GMAIL API ==================
def get_creds():
    creds = None
    if os.path.exists(TOKEN_PATH):
        creds = Credentials.from_authorized_user_file(TOKEN_PATH, SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(CREDENTIALS_PATH, SCOPES)
            creds = flow.run_local_server(port=0)
        token_dir = os.path.dirname(TOKEN_PATH)
        if token_dir:
            ensure_dir(token_dir)
        with open(TOKEN_PATH, "w") as f:
            f.write(creds.to_json())
    return creds

def get_email_subject(service, msg_id):
    try:
        msg = service.users().messages().get(
            userId="me",
            id=msg_id,
            format="metadata"
        ).execute(num_retries=GMAIL_RETRIES)
        headers = msg.get("payload", {}).get("headers", [])
        subject = ""
        for h in headers:
            if h["name"].lower() == "subject":
                subject = h["value"]
                break
        if not subject:
            return "NO_SUBJECT"
        decoded_parts = decode_header(subject)
        decoded_string = ""
        for part, encoding in decoded_parts:
            if isinstance(part, bytes):
                decoded_string += part.decode(encoding or "utf-8", errors="ignore")
            else:
                decoded_string += part
        return decoded_string.strip()
    except Exception as e:
        return f"ERROR_SUBJECT: {e}"

def extract_email_text(html):
    try:
        soup = BeautifulSoup(html, "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        text = soup.get_text(separator="\n")
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        return "\n".join(lines)
    except:
        return "CANNOT_PARSE_EMAIL_CONTENT"

def get_label_id(service, label_name):
    labels = service.users().labels().list(userId="me").execute(num_retries=GMAIL_RETRIES).get("labels", [])
    for l in labels:
        if l["name"].lower() == label_name.lower():
            return l["id"]
    return None

# ================== TẢI PDF ĐÍNH KÈM GMAIL ĐA LUỒNG ==================
def download_single_pdf(msg_id, creds, parts=None):
    """parts: danh sách file đính kèm lấy sẵn khi quét -> tải thẳng bằng attachments.get(), không đọc lại email.
    parts=None (vd. app.py gọi) -> đọc email bằng format="full" như cũ."""
    try:
        local_service = get_local_service(creds)      # dùng lại service của luồng, không build() lại mỗi email
        payload = _email_payload(local_service, msg_id, parts)

        def extract_pdf(payload):
            results = []
            if payload.get("parts"):
                for p in payload["parts"]:
                    results.extend(extract_pdf(p))
            else:
                if is_pdf_part(payload):
                    results.append(payload)
            return results

        pdf_parts = extract_pdf(payload)
        downloaded = 0
        pdf_names = []

        for i, part in enumerate(pdf_parts):
            attachment_id = part["body"].get("attachmentId")
            if not attachment_id:
                continue

            attachment = local_service.users().messages().attachments().get(
                userId="me", messageId=msg_id, id=attachment_id).execute(num_retries=GMAIL_RETRIES)
            data = base64.urlsafe_b64decode(attachment["data"])

            filename = part.get("filename") or f"{msg_id}_{i}.pdf"
            filename = get_unique_filename(DOWNLOAD_DIR, filename)
            path = os.path.join(DOWNLOAD_DIR, filename)

            with open(path, "wb") as f:
                f.write(data)
            pdf_names.append(filename)
            saved_count = _finalize_download(path, filename, "GMAIL_PDF", msg_id)
            downloaded += saved_count

        save_xml_attachments(local_service, msg_id, payload)   # --keep-xml: lưu thêm XML đính kèm (tự bắt lỗi)
        _xml_after_email(msg_id, "PDF đính kèm", pdf_names)
        return downloaded
    except Exception as e:
        if parts is not None:
            # Thông tin lấy lúc quét không dùng được -> đọc lại email như cách cũ (file trùng đã tải sẽ bị _dedupe_pdf bỏ)
            print(f"[WARN] GMAIL_PDF | EMAIL={msg_id} | Tải theo thông tin lúc quét lỗi ({e}) -> đọc lại email")
            return download_single_pdf(msg_id, creds)
        print(f"[ERROR] GMAIL_PDF | EMAIL={msg_id} | ERR={e}")
        return 0

def download_pdfs_parallel(pdf_email_ids, creds, parts_by_id=None):
    total_downloaded = 0
    successful_ids = set()
    parts_by_id = parts_by_id or {}
    with concurrent.futures.ThreadPoolExecutor(max_workers=SCAN_WORKERS) as executor:
        future_to_msg_id = {executor.submit(download_single_pdf, msg_id, creds, parts_by_id.get(msg_id)): msg_id
                            for msg_id in pdf_email_ids}
        for future in concurrent.futures.as_completed(future_to_msg_id):
            downloaded = future.result()
            total_downloaded += downloaded
            if downloaded:
                successful_ids.add(future_to_msg_id[future])
    return total_downloaded, successful_ids

# ================== EXTRACTORS (TÌM LINK) ==================
def extract_easyinvoice_link(html):
    soup = BeautifulSoup(html, "html.parser")
    for a in soup.find_all("a"):
        text = a.get_text(strip=True).lower()
        if text == "link":
            next_text = ""
            for sibling in a.next_siblings:
                if isinstance(sibling, str):
                    next_text += sibling.lower()
                else:
                    next_text += sibling.get_text(strip=True).lower()

                if len(next_text) > 40:
                    break

            if "tải nhanh" in next_text or "dạng pdf" in next_text:
                href = a.get("href")
                if href:
                    return href

    links = soup.find_all("a", string=lambda t: t and t.strip().lower() == "link")
    if len(links) >= 2: return links[1].get("href")
    if len(links) == 1: return links[0].get("href")
    return None

def extract_meinvoice_link(html):
    soup = BeautifulSoup(html, "html.parser")
    for a in soup.find_all("a"):
        text = a.get_text(strip=True).lower()
        if "tra cứu" in text:
            href = a.get("href")
            if href and "meinvoice.vn" in href: return href
    return None

def extract_pvoil_link(html):
    soup = BeautifulSoup(html, "html.parser")
    for a in soup.find_all("a"):
        text = a.get_text(strip=True).lower()
        if "tại đây" in text or "here" in text:
            href = a.get("href")
            if href and "hoadon.pvoil.vn" in href: return href
    return None

def extract_smartsign_link(html):
    soup = BeautifulSoup(html, "html.parser")
    smartsign_links = []
    for a in soup.find_all("a"):
        href = a.get("href")
        if not href or "tracuuhd.smartsign.com.vn" not in href:
            continue
        text = a.get_text(strip=True).lower()
        if text == "link":
            smartsign_links.append(href)
    for href in smartsign_links:
        if "export=pdf" in href.lower():
            return href
    return smartsign_links[0] if smartsign_links else None

def extract_fast_link(html):
    soup = BeautifulSoup(html, "html.parser")
    for a in soup.find_all("a"):
        text = a.get_text(strip=True).lower()
        if "nhấn vào đây" in text or "click here" in text:
            parent_text = a.parent.get_text(strip=True).lower()
            if "pdf file" in parent_text:
                href = a.get("href")
                if href: return href
    return None

def extract_ehoadon_link(html):
    soup = BeautifulSoup(html, "html.parser")
    for a in soup.find_all("a"):
        href = a.get("href")
        text = a.get_text(strip=True).upper()
        if href and "ehoadon.vn" in href:
            if "hdpb6a" in href: continue
            if "TRACUU.EHOADON.VN" in text or text.isalnum(): return href
    return None

def extract_vnpt_link(html):
    soup = BeautifulSoup(html, "html.parser")
    for a in soup.find_all("a"):
        text = a.get_text(strip=True).lower()
        href = a.get("href", "")
        if "tải pdf" in text or "tai pdf" in text:
            if href: return href
    for a in soup.find_all("a"):
        href = a.get("href", "")
        if href and "vnpt-invoice.com.vn" in href and "pdf" in href.lower(): return href
    for a in soup.find_all("a"):
        href = a.get("href", "")
        if href and "vnpt-invoice.com.vn" in href: return href
    return None

# ================== HELPERS (ĐẶT TÊN FILE) ==================
import re as _re

def _parse_invoice_info(page_text):
    ky_hieu = None
    so_hd   = None
    mst     = None

    m = _re.search(r'K[yý] hi[eệ]u\s*[:\s]\s*([A-Z0-9]{5,10})', page_text, _re.IGNORECASE)
    if m: ky_hieu = m.group(1).strip()

    m = _re.search(r'S[oố]\s*[:\s]\s*(\d+)', page_text, _re.IGNORECASE)
    if m: so_hd = m.group(1).strip()

    m = _re.search(r'(\d{10})', page_text)
    if m: mst = m.group(1).strip()

    return {"kyHieu": ky_hieu, "soHd": so_hd, "mst": mst}

def _build_filename(inv, suggested, supplier_tag):
    ky_hieu = inv.get("kyHieu")
    so_hd   = inv.get("soHd")
    mst     = inv.get("mst") or "0000000000"
    if ky_hieu and so_hd: return f"{mst}_{ky_hieu}_{so_hd}.pdf"
    return suggested or f"{supplier_tag}_invoice.pdf"

def _extract_pdf_text(file_path, max_pages=2):
    if not pdfplumber:
        return ""
    try:
        with pdfplumber.open(file_path) as pdf:
            pages = pdf.pages[:max_pages]
            return "\n".join(page.extract_text() or "" for page in pages)
    except Exception as e:
        print(f"[WARN] Khong doc duoc noi dung PDF de dat ten: {file_path} | ERR={repr(e)}")
        return ""

def _build_filename_from_pdf(file_path, page_inv, suggested, supplier_tag):
    pdf_text = _extract_pdf_text(file_path)
    pdf_inv = _parse_invoice_info(pdf_text) if pdf_text else {}
    final_name = _build_filename(pdf_inv, None, supplier_tag)
    if final_name != f"{supplier_tag}_invoice.pdf":
        page_name = _build_filename(page_inv, suggested, supplier_tag)
        if page_name != final_name:
            print(f"[WARN] SMARTSIGN_FILENAME_MISMATCH | PAGE_NAME={page_name} | PDF_NAME={final_name}")
        return final_name
    return _build_filename(page_inv, suggested, supplier_tag)

def _smartsign_after_click_url(url):
    parsed = urllib.parse.urlparse(url)
    query = urllib.parse.parse_qs(parsed.query)
    code = (query.get("code") or [""])[0].strip()
    if not code:
        return url
    code = urllib.parse.quote(code, safe="")
    return f"https://tracuuhdc.smartsign.com.vn/hddt/?code={code}&xdb=1"

# ================== WORKERS (TẢI FILE) ==================
async def process_easyinvoice(context, url, sem, email_id, supplier):
    async with sem:
        page = await context.new_page()
        try:
            download = None
            try:
                async with page.expect_download(timeout=15000) as download_info:
                    await page.goto(url)
                download = await download_info.value
            except Exception:
                await page.goto(url, timeout=60000, wait_until="networkidle")
                await page.wait_for_timeout(2000)
                await page.wait_for_selector("text=PDF", timeout=15000)
                async with page.expect_download(timeout=60000) as download_info:
                    await page.click("text=PDF", force=True)
                download = await download_info.value

            if download:
                filename = download.suggested_filename
                filename = get_unique_filename(DOWNLOAD_DIR, filename)
                path = os.path.join(DOWNLOAD_DIR, filename)
                await download.save_as(path)
                saved_count = _finalize_download(path, filename, supplier, email_id)
                if saved_count:
                    return saved_count
                raise Exception(f"Không lưu được PDF từ file tải về: {filename}")
            raise Exception("Không bắt được sự kiện tải file.")
        except Exception as e:
            print(f"[ERROR] {supplier.upper()} | EMAIL={email_id} | URL={url} | ERR={repr(e)}")
            raise
        finally:
            await page.close()

async def process_meinvoice(context, url, sem, email_id, supplier):
    async with sem:
        page = await context.new_page()
        try:
            async with page.expect_download() as d:
                await page.goto(url)
            download = await d.value
            filename = get_unique_filename(DOWNLOAD_DIR, download.suggested_filename)
            path = os.path.join(DOWNLOAD_DIR, filename)
            await download.save_as(path)
            saved_count = _finalize_download(path, filename, supplier, email_id)
            if saved_count:
                return saved_count
            raise Exception(f"Không lưu được PDF từ file tải về: {filename}")
        except Exception as e:
            print(f"[ERROR] {supplier.upper()} | EMAIL={email_id} | URL={url} | ERR={repr(e)}")
            raise
        finally:
            await page.close()

async def process_pvoil(context, url, sem, email_id, supplier):
    async with sem:
        page = await context.new_page()
        try:
            download = None
            try:
                async with page.expect_download(timeout=15000) as d:
                    await page.goto(url, timeout=60000)
                download = await d.value
            except:
                download = await page.wait_for_event("download", timeout=15000)
            filename = get_unique_filename(DOWNLOAD_DIR, download.suggested_filename)
            path = os.path.join(DOWNLOAD_DIR, filename)
            await download.save_as(path)
            saved_count = _finalize_download(path, filename, supplier, email_id)
            if saved_count:
                return saved_count
            raise Exception(f"Không lưu được PDF từ file tải về: {filename}")
        except Exception as e:
            print(f"[ERROR] {supplier.upper()} | EMAIL={email_id} | URL={url} | ERR={repr(e)}")
            raise
        finally:
            await page.close()

# ---- Giai đoạn 3: tải XML từ trang tra cứu SmartSign (chỉ khi bật --keep-xml) ----
# Trang xem hóa đơn có nút "Tải file XML" ngay cạnh nút "Tải file PDF" (xác nhận trên trang thật).
SMARTSIGN_SEL_XML = (
    "a:has-text('Tải file XML'), button:has-text('Tải file XML'), input[value*='Tải file XML'], "
    "a:has-text('Tải File XML'), button:has-text('Tải File XML'), input[value*='Tải File XML'], "
    "a:has-text('Tai file XML'), button:has-text('Tai file XML'), "
    "a:has-text('Tải XML'), button:has-text('Tải XML'), input[value*='Tải XML'], "
    "a[href*='export=xml']"
)

async def _smartsign_wait_button(loc, timeout=15000):
    """Chờ nút/link xuất hiện trên trang (trang SmartSign có lúc hiện chậm khi máy chủ bận)."""
    try:
        await loc.wait_for(state="attached", timeout=timeout)
        return True
    except Exception:
        return False

async def _smartsign_page_snippet(page, limit=150):
    try:
        t = await page.inner_text("body", timeout=5000)
    except Exception:
        return "(không đọc được nội dung trang)"
    t = " ".join((t or "").split())
    return t[:limit] if t else "(trang trắng)"

async def _smartsign_save_xml(page, pdf_name, email_id):
    """Bấm "Tải file XML" trên trang SmartSign đang mở, lưu vào thư mục XML (cùng tên với PDF).
    Lỗi chỉ ghi cảnh báo, không ảnh hưởng PDF đã tải."""
    try:
        btn = page.locator(SMARTSIGN_SEL_XML).first
        if not await _smartsign_wait_button(btn, 5000):
            print(f"[WARN] [KEEP-XML] SMARTSIGN | EMAIL={email_id} | Không thấy nút 'Tải file XML' trên trang tra cứu")
            _xml_note(email_id, error="Không thấy nút 'Tải file XML' trên trang SmartSign")
            return ""
        async with page.expect_download(timeout=30000) as d:
            await btn.evaluate("(el) => el.click()")
        download = await d.value
        tmp_path = await download.path()
        with open(tmp_path, "rb") as f:
            data = f.read()
        base = os.path.splitext(pdf_name)[0] + ".xml"
        if data[:2] == b"PK":
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                return "zip" if save_xml_from_zip(z, f"SMARTSIGN:{pdf_name}", email_id) else ""
        return save_invoice_xml(data, base, "SMARTSIGN_XML", email_id)
    except Exception as e:
        print(f"[WARN] [KEEP-XML] SMARTSIGN | EMAIL={email_id} | Không tải được XML (PDF vẫn đã lưu) | ERR={repr(e)[:200]}")
        _xml_note(email_id, error=f"Không tải được XML từ trang SmartSign: {repr(e)[:200]}")
        return ""

async def process_smartsign(context, url, sem, email_id, supplier):
    async with sem:
        page = await context.new_page()
        try:
            download = None
            target_url = _smartsign_after_click_url(url)
            tried_urls = []

            async def open_smartsign_like_email():
                safe_url = html.escape(url, quote=True)
                await page.set_content(
                    "<!doctype html><html><body>"
                    f"<p>Vui lòng truy cập <a id='smartlink' href='{safe_url}'>Link</a> "
                    "để tải hóa đơn định dạng PDF.</p>"
                    "</body></html>"
                )
                tried_urls.append(url)
                try:
                    async with page.expect_navigation(wait_until="domcontentloaded", timeout=60000):
                        await page.click("#smartlink")
                except Exception:
                    if page.url == "about:blank":
                        await page.click("#smartlink")
                try:
                    await page.wait_for_load_state("domcontentloaded", timeout=15000)
                except Exception:
                    pass

            await open_smartsign_like_email()

            if download is None and "login.aspx" in page.url.lower():
                for candidate in [target_url, url]:
                    tried_urls.append(candidate)
                    try:
                        await page.goto(candidate, wait_until="domcontentloaded", timeout=60000)
                    except Exception:
                        await page.goto(candidate, wait_until="domcontentloaded", timeout=60000)
                    if "login.aspx" not in page.url.lower():
                        break

            if "login.aspx" in page.url.lower():
                raise Exception(
                    "SmartSign chuyển về trang đăng nhập, không mở được link hóa đơn: "
                    f"{page.url} | TRIED={tried_urls}"
                )

            if download is None:
                pdf_link = page.locator(
                    "a:has-text('Tải PDF'), "
                    "button:has-text('Tải PDF'), "
                    "input[value*='Tải PDF'], "
                    "a:has-text('Tai PDF'), "
                    "button:has-text('Tai PDF'), "
                    "input[value*='Tai PDF'], "
                    "a:has-text('Tải File PDF'), "
                    "button:has-text('Tải File PDF'), "
                    "input[value*='Tải File PDF'], "
                    "a:has-text('Tải file PDF'), "
                    "button:has-text('Tải file PDF'), "
                    "input[value*='Tải file PDF'], "
                    "a:has-text('Tai file PDF'), "
                    "button:has-text('Tai file PDF'), "
                    "input[value*='Tai file PDF'], "
                    "a[href*='export=pdf'], "
                    "a[href*='.pdf']"
                ).first
                # Chờ nút hiện ra; nếu chưa có (máy chủ chậm/bận) thì nghỉ rồi tải lại trang, tối đa 2 lần
                found = await _smartsign_wait_button(pdf_link, 15000)
                attempt = 0
                while not found and attempt < 2:
                    attempt += 1
                    print(f"[WARN] SMARTSIGN | EMAIL={email_id} | Chưa thấy nút tải PDF, tải lại trang lần {attempt}")
                    await asyncio.sleep(3 * attempt)
                    try:
                        if page.url.lower().startswith("http"):
                            await page.reload(wait_until="domcontentloaded", timeout=60000)
                        else:
                            await page.goto(url, wait_until="domcontentloaded", timeout=60000)
                    except Exception:
                        pass
                    found = await _smartsign_wait_button(pdf_link, 15000)
                if not found:
                    snippet = await _smartsign_page_snippet(page)
                    raise Exception(f"Không tìm thấy nút/link tải PDF trên trang SmartSign. TRANG={snippet}")

            page_text = await page.inner_text("body")
            inv = _parse_invoice_info(page_text)

            if download is None:
                async with page.expect_download(timeout=30000) as d:
                    await pdf_link.evaluate("(el) => el.click()")
                download = await d.value

            suggested = download.suggested_filename
            tmp_name = get_unique_filename(DOWNLOAD_DIR, f"__tmp_smartsign_{email_id}.pdf")
            tmp_path = os.path.join(DOWNLOAD_DIR, tmp_name)
            await download.save_as(tmp_path)
            base_name = _build_filename_from_pdf(tmp_path, inv, suggested, "smartsign")
            filename = get_unique_filename(DOWNLOAD_DIR, base_name)
            path = os.path.join(DOWNLOAD_DIR, filename)
            os.replace(tmp_path, path)
            saved_count = _finalize_download(path, filename, supplier, email_id)
            if saved_count and KEEP_XML:
                await _smartsign_save_xml(page, filename, email_id)   # --keep-xml: lấy thêm XML (tự bắt lỗi)
            if saved_count:
                return saved_count
            raise Exception(f"Không lưu được PDF từ file tải về: {filename}")
        except Exception as e:
            print(f"[ERROR] {supplier.upper()} | EMAIL={email_id} | URL={url} | ERR={repr(e)}")
            raise
        finally:
            await page.close()

async def process_fast(context, url, sem, email_id, supplier):
    async with sem:
        page = await context.new_page()
        try:
            async with page.expect_download(timeout=30000) as download_info:
                try:
                    await page.goto(url)
                except Exception as goto_err:
                    if "Download is starting" not in str(goto_err):
                        raise goto_err

            download = await download_info.value
            filename = get_unique_filename(DOWNLOAD_DIR, f"HoaDon_Fast_{download.suggested_filename}")
            path = os.path.join(DOWNLOAD_DIR, filename)
            await download.save_as(path)
            saved_count = _finalize_download(path, filename, supplier, email_id)
            if saved_count:
                return saved_count
            raise Exception(f"Không lưu được PDF từ file tải về: {filename}")
        except Exception as e:
            print(f"[ERROR] {supplier.upper()} | EMAIL={email_id} | URL={url} | ERR={repr(e)}")
            raise
        finally:
            await page.close()

async def process_ehoadon(context, url, sem, email_id, supplier):
    async with sem:
        page = await context.new_page()
        try:
            ma_tra_cuu = url.rstrip('/').split('/')[-1]
            real_url = f"https://tchd.ehoadon.vn/TCHD?MTC={ma_tra_cuu}"

            filename = get_unique_filename(DOWNLOAD_DIR, f"HoaDon_Bkav_{ma_tra_cuu}.pdf")
            path = os.path.join(DOWNLOAD_DIR, filename)

            await page.goto(real_url)

            frame = page.frame_locator('iframe[name="frameViewInvoice"]')
            object_tag = frame.locator('object[type="application/pdf"]')

            await object_tag.wait_for(state="attached", timeout=15000)
            pdf_relative_url = await object_tag.get_attribute("data")

            if pdf_relative_url:
                pdf_absolute_url = urllib.parse.urljoin(real_url, pdf_relative_url)
                async with page.expect_download(timeout=30000) as download_info:
                    await page.evaluate("""async ([pdf_url, name]) => {
                        const response = await fetch(pdf_url);
                        const blob = await response.blob();
                        const objectUrl = window.URL.createObjectURL(blob);

                        const a = document.createElement('a');
                        a.href = objectUrl;
                        a.download = name;
                        document.body.appendChild(a);
                        a.click();
                        document.body.removeChild(a);

                        setTimeout(() => window.URL.revokeObjectURL(objectUrl), 1000);
                    }""", [pdf_absolute_url, filename])

                download = await download_info.value
                await download.save_as(path)
                saved_count = _finalize_download(path, filename, supplier, email_id)
                if saved_count:
                    return saved_count
                raise Exception(f"Không lưu được PDF từ file tải về: {filename}")
            else:
                raise Exception(f"Không tìm thấy dữ liệu PDF tại (Bkav): {real_url}")
        except Exception as e:
            print(f"[ERROR] {supplier.upper()} | EMAIL={email_id} | URL={url} | ERR={repr(e)}")
            raise
        finally:
            await page.close()

async def process_vnpt(context, url, sem, email_id, supplier):
    async with sem:
        page = await context.new_page()
        try:
            await page.goto("about:blank")
            async with page.expect_download(timeout=30000) as d:
                await page.evaluate("""(url) => {
                    const a = document.createElement('a');
                    a.href = url;
                    a.download = '';
                    document.body.appendChild(a);
                    a.click();
                }""", url)
            download = await d.value
            suggested = download.suggested_filename or "vnpt_invoice.pdf"
            filename = get_unique_filename(DOWNLOAD_DIR, suggested)
            path = os.path.join(DOWNLOAD_DIR, filename)
            await download.save_as(path)
            saved_count = _finalize_download(path, filename, supplier, email_id)
            if saved_count:
                return saved_count
            raise Exception(f"Không lưu được PDF từ file tải về: {filename}")
        except Exception as e:
            print(f"[ERROR] {supplier.upper()} | EMAIL={email_id} | URL={url} | ERR={repr(e)}")
            raise
        finally:
            await page.close()



# ---- Giai đoạn 3 (tiếp): tải XML cho các nhà cung cấp còn lại (chỉ khi bật --keep-xml) ----
# Chạy SAU KHI PDF đã tải xong; mọi lỗi chỉ ghi cảnh báo, không ảnh hưởng PDF / trạng thái email.
#   VNPT        : email có sẵn link "tải XML" (/invoice/getinvoice?token=...)
#   Fast        : email có sẵn link "To download the XML file" (index.aspx?...&type=2...)
#   MeInvoice   : mở trang tra cứu www.meinvoice.vn/tra-cuu/?sc=<mã> rồi bấm "Tải hóa đơn dạng XML"
#   EasyInvoice : mở trang xem hóa đơn (Invoice/ViewFromEmail) rồi bấm nút tải XML
#   PVOil       : mở trang xem hóa đơn (Invoice/ViewFromEmail) rồi bấm nút tải XML
#   Bkav        : mở trang tra cứu tchd.ehoadon.vn rồi bấm nút tải XML
WEB_XML_SUPPLIERS = {"vnpt", "fast", "meinvoice", "easyinvoice", "pvoil", "ehoadon"}
EHOADON_TCHD_URL = "https://tchd.ehoadon.vn/TCHD?MTC={mtc}"
WEB_XML_MAX_MISS = 3              # 1 hãng lỗi XML liên tiếp 3 email -> tạm dừng lấy XML hãng đó trong lần chạy (đỡ chậm)
_WEB_XML_MISS = {}

def _looks_like_invoice_xml(data):
    head = (data or b"")[:3000].lstrip(b"\xef\xbb\xbf").lstrip()
    if not head.startswith(b"<"):
        return False
    low = head[:800].lower()
    if b"<html" in low or b"<!doctype html" in low:
        return False
    return any(tag in data for tag in (b"<HDon", b"<TDiep", b"DLHDon", b"<Invoice", b"<inv:"))

def _email_links(html_text):
    """[(href, chữ trên link, chữ ngay trước link)] của nội dung email."""
    out = []
    try:
        soup = BeautifulSoup(html_text or "", "html.parser")
        for a in soup.find_all("a"):
            href = (a.get("href") or "").strip()
            if not href:
                continue
            before = ""
            for sib in a.previous_siblings:
                before = (sib if isinstance(sib, str) else sib.get_text(" ", strip=True)) + " " + before
                if len(before) > 120:
                    break
            out.append((href, a.get_text(" ", strip=True), before.strip()))
    except Exception:
        pass
    return out

def web_xml_plan(item):
    """Danh sách cách lấy XML cho 1 email: [("get", url) | ("page", url)]."""
    sup = item.get("supplier") or ""
    url = item.get("url") or ""
    links = _email_links(item.get("content") or "")
    plan = []
    if sup == "vnpt":
        for href, text, before in links:
            if "getinvoice" in href.lower() or "xml" in text.lower():
                plan.append(("get", href))
                break
    elif sup == "fast":
        for href, text, before in links:
            if "type=2" in href.lower() or "xml" in before.lower()[-60:]:
                plan.append(("get", href))
                break
    elif sup == "meinvoice":
        q = urllib.parse.parse_qs(urllib.parse.urlparse(url).query)
        code = ((q.get("sc") or q.get("code") or q.get("Code") or [""])[0]).strip()
        if code:
            p = urllib.parse.urlparse(url)
            plan.append(("page", f"{p.scheme}://{p.netloc}/tra-cuu/?sc={urllib.parse.quote(code)}"))
    elif sup in ("easyinvoice", "pvoil"):
        view = ""
        for href, text, before in links:
            if "viewfromemail" in href.lower():
                view = href
                break
        if not view and url:
            view = re.sub(r"(?i)/Invoice/(DownloadInvPdf|getinvoice)", "/Invoice/ViewFromEmail", url)
            if view == url:
                view = ""
        if view:
            plan.append(("page", view))
    elif sup == "ehoadon":
        mtc = url.rstrip("/").split("/")[-1]
        if mtc:
            plan.append(("page", EHOADON_TCHD_URL.format(mtc=mtc)))
    return plan

async def _web_get_bytes(context, url):
    resp = await context.request.get(url, timeout=30000)
    if not resp.ok:
        raise Exception(f"HTTP {resp.status}")
    return await resp.body()

_PICK_XML_JS = r"""() => {
  document.querySelectorAll('[data-kx-pick]').forEach(e => e.removeAttribute('data-kx-pick'));
  const cands = [];
  for (const el of document.querySelectorAll('a,button,input,[onclick],[role=button]')) {
    const t = ((el.innerText || el.value || el.title || '') + '').trim();
    if (t.length > 60) continue;
    const h = (el.getAttribute('href') || '') + ' ' + (el.getAttribute('onclick') || '');
    if (!/xml/i.test(t + ' ' + h)) continue;
    if (/hướng dẫn|huong dan|gửi|email/i.test(t)) continue;
    let score = /xml/i.test(t) ? 2 : 1;
    if (/t[aả]i|download/i.test(t)) score += 2;
    if (/t[ệe]p xml|d[ạa]ng xml|file xml/i.test(t)) score += 2;
    if (['A', 'BUTTON', 'INPUT'].includes(el.tagName)) score += 1;
    cands.push([score, el]);
  }
  if (!cands.length) return false;
  cands.sort((a, b) => b[0] - a[0]);
  cands[0][1].setAttribute('data-kx-pick', '1');
  return true;
}"""

# Chữ trên nút tải XML thật: "Tải hóa đơn dạng XML" (MeInvoice), "Tải tệp XML" (EasyInvoice), "Tải file XML"...
_XML_BTN_RE = re.compile(r"(t[aả]i|download).{0,20}xml", re.I)
# Nút mở menu tải (MeInvoice phải bấm "Tải hóa đơn" thì mới hiện "Tải hóa đơn dạng PDF / XML")
_XML_MENU_RE = re.compile(r"^\s*(t[aả]i\s*h[oó]a\s*đ[oơ]n|t[aả]i\s*v[ềe]|t[aả]i\s*xu[ốo]ng|download)\s*$", re.I)

_XML_SKIP_RE = re.compile(r"h[ưu][ớo]ng\s*d[ẫa]n|g[ửu]i|email|ph[ầa]n\s*m[ềe]m", re.I)

async def _find_visible_text(page, rx):
    for fr in page.frames:
        try:
            loc = fr.get_by_text(rx)
            n = await loc.count()
            for i in range(min(n, 10)):
                el = loc.nth(i)
                if await el.is_visible():
                    txt = (await el.inner_text(timeout=2000)) or ""
                    if len(txt) > 80 or _XML_SKIP_RE.search(txt):
                        continue                             # bỏ "Hướng dẫn tải XML", khung chứa cả trang...
                    return el
        except Exception:
            continue
    return None

async def _catch_file(page, context, action, timeout=25):
    """Bấm rồi chờ file: bắt sự kiện tải file, hoặc trang/tab mới trả thẳng nội dung XML/ZIP."""
    loop = asyncio.get_running_loop()
    fut = loop.create_future()
    def done(kind, obj):
        if not fut.done():
            fut.set_result((kind, obj))
    def on_dl(d):
        done("dl", d)
    def on_resp(r):
        try:
            h = r.headers
            ct = (h.get("content-type") or "").lower()
            cd = (h.get("content-disposition") or "").lower()
            if "attachment" in cd:
                return                                       # file tải về -> chờ sự kiện download
            if ("html" not in ct and ("xml" in ct or "zip" in ct)) or ".xml" in cd or ".zip" in cd:
                done("resp", r)
        except Exception:
            pass
    async def grab_popup(p):                             # nút mở tab mới hiển thị thẳng XML -> lấy lại theo địa chỉ tab
        try:
            await p.wait_for_load_state("domcontentloaded", timeout=timeout * 1000)
            if p.url.lower().startswith("http"):
                data = await _web_get_bytes(context, p.url)
                if data[:2] == b"PK" or _looks_like_invoice_xml(data):
                    done("bytes", data)
        except Exception:
            pass
    tasks = []
    def hook(p):
        p.on("download", on_dl)
        p.on("response", on_resp)
    def on_page(p):
        hook(p)
        tasks.append(asyncio.ensure_future(grab_popup(p)))
    hook(page)
    context.on("page", on_page)
    try:
        await action()
        kind, obj = await asyncio.wait_for(fut, timeout)
        if kind == "dl":
            with open(await obj.path(), "rb") as f:
                return f.read()
        if kind == "bytes":
            return obj
        return await obj.body()
    finally:
        for ev, fn in (("download", on_dl), ("response", on_resp)):
            try:
                page.remove_listener(ev, fn)
            except Exception:
                pass
        try:
            context.remove_listener("page", on_page)
        except Exception:
            pass
        for t in tasks:
            if not t.done():
                t.cancel()

async def _web_click_xml(context, url):
    """Mở trang xem/tra cứu hóa đơn, tìm nút tải XML (mở menu "Tải hóa đơn" nếu cần, cả trong iframe),
    bấm như người dùng và trả về nội dung file tải về."""
    page = await context.new_page()
    try:
        try:
            await page.goto(url, wait_until="domcontentloaded", timeout=60000)
        except Exception as e:
            if "Download is starting" not in str(e):
                raise
        try:
            await page.wait_for_load_state("networkidle", timeout=10000)
        except Exception:
            pass
        btn, menu_opened = None, 0
        for _ in range(8):                                   # chờ tối đa ~16 giây cho nút hiện ra
            btn = await _find_visible_text(page, _XML_BTN_RE)
            if btn:
                break
            if menu_opened < 2:                              # thử mở menu "Tải hóa đơn" (tối đa 2 lần)
                menu = await _find_visible_text(page, _XML_MENU_RE)
                if menu:
                    try:
                        await menu.click(timeout=5000)
                        menu_opened += 1
                        await page.wait_for_timeout(800)
                        continue
                    except Exception:
                        pass
            await page.wait_for_timeout(2000)
        errors = []
        if btn:                                              # cách 1: bấm chuột thật vào nút đang hiện
            try:
                return await _catch_file(page, context, lambda: btn.click(timeout=10000))
            except Exception as e:
                errors.append(f"bấm nút: {repr(e)[:120]}")
        frame = None                                         # cách 2: tìm nút XML (kể cả đang ẩn) rồi bấm bằng JS
        for fr in page.frames:
            try:
                if await fr.evaluate(_PICK_XML_JS):
                    frame = fr
                    break
            except Exception:
                continue
        if frame:
            try:
                return await _catch_file(page, context,
                                         lambda: frame.locator("[data-kx-pick]").first.evaluate("(el) => el.click()"))
            except Exception as e:
                errors.append(f"bấm JS: {repr(e)[:120]}")
        snippet = await _smartsign_page_snippet(page)
        if not btn and not frame:
            raise Exception(f"Không thấy nút tải XML trên trang. TRANG={snippet}")
        raise Exception(f"Đã bấm nút XML nhưng không nhận được file ({'; '.join(errors)}). TRANG={snippet}")
    finally:
        await page.close()

def _xml_name_from_content(data, fallback):
    try:
        head = data[:20000].decode("utf-8", "ignore")
        kh = re.search(r"<KHHDon>\s*([^<\s]+)\s*</KHHDon>", head)
        so = re.search(r"<SHDon>\s*(\d+)\s*</SHDon>", head)
        if kh and so:
            return f"{kh.group(1)}_{so.group(1)}.xml"
    except Exception:
        pass
    return fallback

def _data_desc(data):
    data = data or b""
    head = data[:40].decode("utf-8", "ignore").replace("\n", " ").replace("\r", " ").strip()
    return f"{len(data)} byte, đầu file: {head[:40]!r}"

def _save_web_xml(data, base, supplier, email_id, url, has_pdf=True):
    if (data or b"")[:2] == b"PK":
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return save_xml_from_zip(z, f"{supplier.upper()}:{base}", email_id) > 0
    if _looks_like_invoice_xml(data):
        if not has_pdf:                                     # chưa có PDF -> đặt tên theo Ký hiệu + Số hóa đơn
            base = _xml_name_from_content(data, base)
        return bool(save_invoice_xml(data, base, f"{supplier.upper()}_XML", email_id))
    return False

async def fetch_web_xml(context, item, sem):
    """--keep-xml: sau khi PDF qua link đã tải xong, lấy thêm XML từ link trong email / trang tra cứu."""
    eid = item.get("email_id") or ""
    sup = item.get("supplier") or ""
    try:
        with _XML_LOCK:
            if (_XML_STATUS.get(eid) or {}).get("files"):
                return                                      # email đã có XML đính kèm
        with _EMAIL_FILES_LOCK:
            pdfs = [n for n in _EMAIL_FILES.get(eid, []) if str(n).lower().endswith(".pdf")]
        base = (os.path.splitext(pdfs[0])[0] if pdfs else f"{sup}_{eid}") + ".xml"
        if _WEB_XML_MISS.get(sup, 0) >= WEB_XML_MAX_MISS:
            return                                          # hãng này đang tạm dừng lấy XML (đã báo 1 lần)
        plan = web_xml_plan(item)
        if not plan:
            print(f"[KEEP-XML] {sup.upper()} | EMAIL={eid} | Email không có link tải XML")
            return
        errors = []
        async with sem:
            for how, u in plan:
                try:
                    data = await (_web_get_bytes(context, u) if how == "get" else _web_click_xml(context, u))
                    if _save_web_xml(data, base, sup, eid, u, bool(pdfs)):
                        _WEB_XML_MISS[sup] = 0
                        return
                    errors.append(f"{'link' if how == 'get' else 'trang'}: nội dung tải về không phải XML hóa đơn ({_data_desc(data)})")
                except Exception as e:
                    errors.append(f"{'link' if how == 'get' else 'trang'}: {repr(e)[:200]}")
        msg = " ; ".join(errors)
        print(f"[WARN] [KEEP-XML] {sup.upper()} | EMAIL={eid} | Không tải được XML" + (" (PDF vẫn đã lưu)" if pdfs else "") + f" | ERR={msg[:300]}")
        _xml_note(eid, error=f"Không tải được XML từ {sup}: {msg[:250]}")
        _WEB_XML_MISS[sup] = _WEB_XML_MISS.get(sup, 0) + 1
        if _WEB_XML_MISS[sup] == WEB_XML_MAX_MISS:
            print(f"[WARN] [KEEP-XML] {sup.upper()} | Lỗi XML {WEB_XML_MAX_MISS} email liên tiếp -> tạm dừng lấy XML "
                  f"của hãng này trong lần chạy này (PDF vẫn tải bình thường)")
    except Exception as e:
        print(f"[WARN] [KEEP-XML] {sup.upper()} | EMAIL={eid} | Lỗi khi lấy XML (PDF vẫn đã lưu) | ERR={repr(e)[:200]}")
        _xml_note(eid, error=f"Lỗi khi lấy XML từ {sup}: {repr(e)[:200]}")


# ================== PETROLIMEX (XỬ LÝ SAU CÙNG) ==================
# Gộp từ petro.py. Email Petro không có file PDF/link để tải, chỉ cần trích xuất
# "Số hóa đơn" + "Mã tra cứu" từ nội dung mail rồi xuất Excel.
def petro_get_body_text(payload):
    """Lấy nội dung thân mail (HTML hoặc text) - trả về (nội dung, mime)."""
    if payload.get("parts"):
        for part in payload["parts"]:
            body, mime = petro_get_body_text(part)
            if body:  # bỏ qua phần không có nội dung (vd. đính kèm) và tìm tiếp phần sau
                return body, mime
    else:
        mime = payload.get("mimeType", "")
        data = payload.get("body", {}).get("data", "")
        if data and mime in ("text/html", "text/plain"):
            decoded = base64.urlsafe_b64decode(data).decode("utf-8", errors="ignore")
            return decoded, mime
    return None, None

def petro_clean_code(value, keep_trailing_star=False):
    if not value:
        return None
    value = str(value).strip()
    value = re.sub(r"^[\s:*#\-]+", "", value)
    value = re.sub(r"[\s:#\-]+$", "", value)
    value = re.sub(r"\s+", "", value)
    if keep_trailing_star:
        value = value.rstrip("*") + "*"
    else:
        value = value.rstrip("*")
    return value or None

def petro_extract_fields(body_text, mime):
    """Trích xuất Số hóa đơn và Mã tra cứu từ nội dung mail."""
    if mime == "text/html":
        soup = BeautifulSoup(body_text, "html.parser")
        text = soup.get_text(separator="\n")
    else:
        text = body_text

    so_hoa_don = None
    ma_tra_cuu = None

    # Mail Petrolimex chuyển tiếp thường bọc giá trị bằng dấu '*', ví dụ:
    # "Mã tra cứu: * 5CTEFB8AG* *".
    patterns_so = [
        r"S[oố]\s*h[oó]a\s*[dđ][oơ]n\s*[:\-]?\s*[*\s]*([A-Z0-9]+)",
        r"Invoice\s*No\.?\s*[:\-]?\s*[*\s]*([A-Z0-9]+)",
    ]
    patterns_ma = [
        r"M[aã]\s*tra\s*c[uứ]u\s*[:\-]?\s*[*\s]*([A-Z0-9]+)\*?",
        r"Lookup\s*Code\s*[:\-]?\s*[*\s]*([A-Z0-9]+)\*?",
    ]

    for pat in patterns_so:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            so_hoa_don = petro_clean_code(m.group(1))
            break

    for pat in patterns_ma:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            ma_tra_cuu = petro_clean_code(m.group(1), keep_trailing_star=True)
            break

    return so_hoa_don, ma_tra_cuu

def petro_header(msg, name):
    for h in msg.get("payload", {}).get("headers", []):
        if h["name"].lower() == name:
            return h["value"]
    return ""

def petro_fetch_record(service, msg_id):
    """Đọc một email Petro và trích xuất dữ liệu (cùng cách làm với petro.py)."""
    msg = service.users().messages().get(userId="me", id=msg_id, format="full").execute(num_retries=GMAIL_RETRIES)
    body_text, mime = petro_get_body_text(msg["payload"])
    so_hoa_don, ma_tra_cuu = (None, None)
    if body_text:
        so_hoa_don, ma_tra_cuu = petro_extract_fields(body_text, mime)
    return {
        "email_id": msg_id,
        "date": petro_header(msg, "date"),
        "sender": petro_header(msg, "from"),
        "subject": petro_header(msg, "subject"),
        "so_hoa_don": so_hoa_don or "",
        "ma_tra_cuu": ma_tra_cuu or "",
    }

def process_petro_emails(gmail_service, state, petro_items, forced_ids, record):
    """
    Xử lý các email Petro (chạy SAU CÙNG, khi các email khác đã xong).
      - mail thuộc label Petro (forced) luôn được coi là Petro
      - mail chỉ khớp từ khóa nhưng không trích xuất được gì -> coi là không phải hóa đơn Petro (SKIPPED)
    """
    stats = {"ok": 0, "missing": 0, "not_petro": 0, "error": 0, "saved": 0}
    for item in petro_items:
        eid = item["id"]
        meta = dict(supplier="petrolimex", subject=item.get("subject"), mail_date=item.get("mail_date"))
        try:
            rec = petro_fetch_record(gmail_service, eid)
        except Exception as e:
            stats["error"] += 1
            print(f"[ERROR] PETRO | EMAIL={eid} | ERR={repr(e)}")
            record(eid, FAILED, reason=f"PETRO_FETCH_ERROR: {repr(e)}"[:300], **meta)
            continue

        so, ma = rec["so_hoa_don"], rec["ma_tra_cuu"]
        if not (so or ma) and eid not in forced_ids:
            stats["not_petro"] += 1
            record(eid, SKIPPED, reason="PETRO_KEYWORD_BUT_NO_INVOICE_FIELDS", **meta)
            continue

        state.save_petro(eid, item.get("mail_date"), rec["date"], rec["sender"], rec["subject"], so, ma, commit=False)
        stats["saved"] += 1
        if so and ma:
            stats["ok"] += 1
            print(f"[SUCCESS] PETRO | EMAIL={eid} | SO_HOA_DON={so} | MA_TRA_CUU={ma}")
            record(eid, DONE, reason="PETRO_EXTRACTED", **meta)
        else:
            stats["missing"] += 1
            thieu = " và ".join(n for n, v in (("số hóa đơn", so), ("mã tra cứu", ma)) if not v)
            print(f"[WARN] PETRO | EMAIL={eid} | THIẾU {thieu} | SUBJECT={rec['subject']}")
            print(f"[WARN] PETRO | OPEN=https://mail.google.com/mail/u/0/#all/{eid}")
            record(eid, FAILED, reason=f"PETRO_MISSING_DATA: thiếu {thieu}", **meta)
    state.commit()
    return stats

# ---------- Petro: FILE TỔNG các hóa đơn đã tải (đọc để lọc trùng, ghi thêm hóa đơn mới tải) ----------
PETRO_MASTER_HEADERS = ["STT", "Ngày nhận", "Người gửi", "Tiêu đề email", "Số hóa đơn", "Mã tra cứu",
                        "File PDF", "Ngày tải", "Email ID"]
# Từ khóa nhận diện cột (không phân biệt hoa/thường, có dấu hay không) -> file tổng của bạn có thể đặt cột theo thứ tự khác
PETRO_MASTER_COLS = {
    "stt": ["stt"],
    "mail_date": ["ngay nhan"],
    "sender": ["nguoi gui"],
    "subject": ["tieu de"],
    "so_hoa_don": ["so hoa don", "invoice no"],
    "ma_tra_cuu": ["ma tra cuu", "lookup"],
    "pdf_file": ["file pdf", "ten file"],
    "down_date": ["ngay tai"],
    "email_id": ["email id", "id email"],
}
_master_backed_up = set()

def _ascii_lower(text):
    t = unicodedata.normalize("NFD", str(text if text is not None else ""))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return t.replace("đ", "d").replace("Đ", "D").lower().strip()

def _petro_master_layout(ws):
    """Tìm dòng tiêu đề và vị trí các cột. Trả về (dòng tiêu đề, {khóa: số cột}) hoặc None nếu không có cột 'Mã tra cứu'."""
    max_col = ws.max_column or 0
    for r in range(1, min(ws.max_row or 0, 15) + 1):
        heads = {}
        for c in range(1, max_col + 1):
            h = _ascii_lower(ws.cell(r, c).value)
            if h:
                heads[c] = h
        cols = {}
        for key, kws in PETRO_MASTER_COLS.items():
            for c, h in heads.items():
                if c in cols.values():
                    continue
                if (h == "stt") if key == "stt" else any(k in h for k in kws):
                    cols[key] = c
                    break
        if "ma_tra_cuu" in cols:
            return r, cols
    return None

def petro_master_load(path):
    """
    Đọc file tổng -> {mã tra cứu chuẩn hóa: {số hóa đơn chuẩn hóa}}.
    Trả về {} nếu chưa có file, None nếu file không đọc được / không nhận ra cột 'Mã tra cứu'.
    """
    if not path or not os.path.exists(path):
        return {}
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        print(f"[WARN] Không đọc được file tổng Petro {path} (đang mở/hỏng?) | ERR={repr(e)}")
        return None
    keys, valid_sheets = {}, 0
    for ws in wb.worksheets:
        layout = _petro_master_layout(ws)
        if not layout:
            continue
        valid_sheets += 1
        hdr, cols = layout
        for r in range(hdr + 1, (ws.max_row or 0) + 1):
            code = norm_code(ws.cell(r, cols["ma_tra_cuu"]).value)
            if not code:
                continue
            so = norm_so(ws.cell(r, cols["so_hoa_don"]).value) if "so_hoa_don" in cols else ""
            keys.setdefault(code, set()).add(so)
    if not valid_sheets:
        print(f"[WARN] File tổng Petro {path}: không tìm thấy cột 'Mã tra cứu' -> không dùng để lọc trùng và không ghi vào file này.")
        return None
    return keys

def petro_master_append(path, rows):
    """
    Thêm các hóa đơn đã tải vào cuối file tổng.
      - File đã có: giữ nguyên dữ liệu/định dạng cũ, chỉ điền vào các cột nhận diện được; sao lưu file gốc thành *.bak.xlsx (1 lần/lần chạy).
      - File chưa có: tạo mới với đầy đủ cột.
    """
    import copy
    import shutil
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter, range_boundaries

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    existed = os.path.exists(path)

    if existed:
        wb = openpyxl.load_workbook(path)
        target = None
        for ws_ in wb.worksheets:
            layout = _petro_master_layout(ws_)
            if layout:
                target = (ws_, layout[0], layout[1])
                break
        if not target:
            raise ValueError("Không nhận ra cột 'Mã tra cứu' trong file tổng")
        ws, hdr, cols = target
        if path not in _master_backed_up:
            root, ext = os.path.splitext(path)
            shutil.copy2(path, f"{root}.bak{ext}")        # vd. ..._tong_hop_down.bak.xlsx - mở trực tiếp bằng Excel được
            _master_backed_up.add(path)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tổng hợp Petrolimex"
        for i, h in enumerate(PETRO_MASTER_HEADERS, 1):
            cell = ws.cell(row=1, column=i, value=h)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill("solid", start_color="1F4E79")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        for col_letter, width in zip("ABCDEFGHI", (6, 22, 32, 42, 16, 18, 42, 20, 20)):
            ws.column_dimensions[col_letter].width = width
        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"
        hdr, cols = _petro_master_layout(ws)
        folder = os.path.dirname(path)
        if folder:
            os.makedirs(folder, exist_ok=True)

    # Dòng dữ liệu cuối cùng (bỏ qua các dòng trống chỉ có định dạng)
    check_cols = [cols["ma_tra_cuu"]] + ([cols["so_hoa_don"]] if "so_hoa_don" in cols else [])
    last = hdr
    for r in range(hdr + 1, (ws.max_row or 0) + 1):
        if any(ws.cell(r, c).value not in (None, "") for c in check_cols):
            last = r

    base = last - hdr
    if "stt" in cols and last > hdr:
        try:
            base = int(ws.cell(last, cols["stt"]).value)
        except (TypeError, ValueError):
            pass

    max_col = max(ws.max_column or 0, max(cols.values()))
    for i, row in enumerate(rows, 1):
        r = last + i
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            if last > hdr:
                cell._style = copy.copy(ws.cell(last, c)._style)     # giữ định dạng của dòng trước
            else:
                cell.font = Font(name="Arial", size=10)
                cell.border = border
                cell.alignment = Alignment(horizontal="left", vertical="center")
        for key, col in cols.items():
            cell = ws.cell(r, col)
            if key == "stt":
                cell.value = base + i
            elif key in ("so_hoa_don", "ma_tra_cuu"):
                cell.value = str(row.get(key) or "")
                cell.number_format = "@"                             # giữ số 0 đầu
            else:
                cell.value = row.get(key) or ""

    new_last = last + len(rows)
    ref = ws.auto_filter.ref
    if ref:
        c1, r1, c2, r2 = range_boundaries(ref)
        ws.auto_filter.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{max(r2, new_last)}"
    elif not existed:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(PETRO_MASTER_HEADERS))}{new_last}"
    wb.save(path)
    return len(rows)

def _petro_dedupe(state, verbose=True):
    """Lọc trùng hóa đơn Petro (trong sổ + so với file tổng). Trả về {'total', 'new'}."""
    keys = None
    if PETRO_MASTER:
        keys = petro_master_load(PETRO_MASTER)
        if keys is None:
            keys = {}                                                # không đọc được file tổng -> chỉ lọc trùng trong sổ
        elif verbose and keys:
            print(f"[PETRO] File tổng: {PETRO_MASTER} ({sum(len(v) for v in keys.values())} hóa đơn đã có)")
    res = state.petro_mark_duplicates(keys)
    if verbose:
        extra = f" ({res['new']} email mới phát hiện)" if res["new"] else ""
        print(f"[PETRO] Lọc trùng (mã tra cứu + số hóa đơn): {res['total']} email trùng -> không tải lại{extra}")
    return res

def _petro_sync_master(state):
    """Ghi các hóa đơn đã tải xong (chưa có trong file tổng) vào file tổng. Trả về số hóa đơn đã thêm."""
    if not PETRO_MASTER:
        return 0
    rows = state.petro_unsynced_master_rows()
    if not rows:
        return 0
    keys = petro_master_load(PETRO_MASTER)
    if keys is None:
        print("[WARN] Không cập nhật được file tổng Petro (xem cảnh báo ở trên). Các hóa đơn đã tải vẫn được nhớ, sẽ ghi vào lần chạy sau.")
        return 0

    to_add, known_ids, seen = [], [], set()
    for r in rows:
        code_n, so_n = norm_code(r["ma_tra_cuu"]), norm_so(r["so_hoa_don"])
        if in_master(keys, code_n, so_n) or (code_n, so_n) in seen:
            known_ids.append(r["message_id"])                        # đã có sẵn trong file tổng -> không thêm lại
            continue
        seen.add((code_n, so_n))
        to_add.append(r)

    try:
        if to_add:
            petro_master_append(PETRO_MASTER, [{
                "mail_date": r["header_date"] or r["mail_date"] or "",
                "sender": r["sender"] or "",
                "subject": r["subject"] or "",
                "so_hoa_don": r["so_hoa_don"] or "",
                "ma_tra_cuu": r["ma_tra_cuu"],
                "pdf_file": r.get("pdf_file") or "Đã tải (thủ công)",
                "down_date": (r.get("pdf_updated_at") or "").replace("T", " ")[:19],
                "email_id": r["message_id"],
            } for r in to_add])
    except PermissionError:
        print(f"[WARN] Không ghi được file tổng {PETRO_MASTER} (đang mở trong Excel?). "
              "Đóng file rồi chạy lại (ví dụ --export-petro) để ghi bổ sung.")
        return 0
    except Exception as e:
        print(f"[WARN] Không ghi được file tổng Petro | ERR={repr(e)}")
        return 0

    state.set_master_synced([r["message_id"] for r in rows])
    extra = f" ({len(known_ids)} hóa đơn đã có sẵn nên không thêm lại)" if known_ids else ""
    print(f"[MASTER] Đã thêm {len(to_add)} hóa đơn vào file tổng: {PETRO_MASTER}{extra}")
    return len(to_add)

def export_petro_excel(data, output_path, source_text="", dup_count=0):
    """Xuất Excel Petro (định dạng giống petro.py)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hóa Đơn Petrolimex"

    # --- Styles ---
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", start_color="1F4E79")
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    alt_fill = PatternFill("solid", start_color="EBF3FB")
    ok_font = Font(name="Arial", color="375623", size=10)
    missing_font = Font(name="Arial", color="9C0006", size=10)
    missing_fill = PatternFill("solid", start_color="FFC7CE")

    # --- Column widths ---
    col_widths = {
        "A": 5,   # STT
        "B": 22,  # Ngày
        "C": 35,  # Người gửi
        "D": 45,  # Tiêu đề
        "E": 18,  # Số hóa đơn
        "F": 18,  # Mã tra cứu
        "G": 12,  # Trạng thái
        "H": 42,  # File PDF
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 30

    # --- Header row ---
    headers = ["STT", "Ngày nhận", "Người gửi", "Tiêu đề email", "Số hóa đơn", "Mã tra cứu", "Trạng thái", "File PDF"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    # --- Data rows ---
    for row_idx, record in enumerate(data, 2):
        row_fill = alt_fill if row_idx % 2 == 0 else None
        status = "✓ OK" if record["so_hoa_don"] and record["ma_tra_cuu"] else "⚠ Thiếu dữ liệu"

        values = [
            row_idx - 1,
            record["date"],
            record["sender"],
            record["subject"],
            record["so_hoa_don"],
            record["ma_tra_cuu"],
            status,
            record.get("pdf_file", ""),
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = left if col_idx > 1 else center
            cell.border = border

            if row_fill and col_idx not in (5, 6, 7, 8):
                cell.fill = row_fill

        # Color-code Số hóa đơn & Mã tra cứu
        for col_idx in (5, 6):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = center
            if cell.value:
                cell.font = ok_font
                cell.fill = PatternFill("solid", start_color="E2EFDA")
            else:
                cell.font = missing_font
                cell.fill = missing_fill

        # Status cell
        status_cell = ws.cell(row=row_idx, column=7)
        status_cell.alignment = center
        if status.startswith("✓"):
            status_cell.font = Font(name="Arial", color="375623", bold=True, size=10)
            status_cell.fill = PatternFill("solid", start_color="E2EFDA")
        else:
            status_cell.font = Font(name="Arial", color="9C0006", bold=True, size=10)
            status_cell.fill = missing_fill

        # File PDF: xanh = đã tải, đỏ = chưa tải
        pdf_cell = ws.cell(row=row_idx, column=8)
        if pdf_cell.value == "Chưa tải":
            pdf_cell.font = missing_font
            pdf_cell.fill = missing_fill
        elif pdf_cell.value and pdf_cell.value != "—":
            pdf_cell.font = ok_font
            pdf_cell.fill = PatternFill("solid", start_color="E2EFDA")

    # --- Summary row ---
    total_rows = len(data)
    ok_count = sum(1 for r in data if r["so_hoa_don"] and r["ma_tra_cuu"])
    missing_count = total_rows - ok_count

    summary_row = total_rows + 3
    ws.cell(row=summary_row, column=1, value="Tổng số email:").font = Font(name="Arial", bold=True, size=10)
    ws.cell(row=summary_row, column=2, value=total_rows).font = Font(name="Arial", size=10)
    ws.cell(row=summary_row + 1, column=1, value="Trích xuất thành công:").font = Font(name="Arial", bold=True, color="375623", size=10)
    ws.cell(row=summary_row + 1, column=2, value=ok_count).font = Font(name="Arial", color="375623", size=10)
    ws.cell(row=summary_row + 2, column=1, value="Thiếu dữ liệu:").font = Font(name="Arial", bold=True, color="9C0006", size=10)
    ws.cell(row=summary_row + 2, column=2, value=missing_count).font = Font(name="Arial", color="9C0006", size=10)
    pdf_done = sum(1 for r in data if r.get("pdf_file") not in (None, "", "—", "Chưa tải"))
    ws.cell(row=summary_row + 3, column=1, value="Đã tải PDF:").font = Font(name="Arial", bold=True, size=10)
    ws.cell(row=summary_row + 3, column=2, value=pdf_done).font = Font(name="Arial", size=10)
    ws.cell(row=summary_row + 4, column=1, value="Email trùng đã lọc:").font = Font(name="Arial", bold=True, size=10)
    ws.cell(row=summary_row + 4, column=2, value=dup_count).font = Font(name="Arial", size=10)

    # --- Freeze header ---
    ws.freeze_panes = "A2"

    # --- Auto-filter ---
    ws.auto_filter.ref = f"A1:H{total_rows + 1}"

    # --- Metadata sheet ---
    meta = wb.create_sheet("Info")
    meta["A1"] = "Tạo lúc:"
    meta["B1"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    meta["A2"] = "Nguồn quét:"
    meta["B2"] = source_text
    meta["A3"] = "Tổng email:"
    meta["B3"] = total_rows

    folder = os.path.dirname(output_path)
    if folder:
        os.makedirs(folder, exist_ok=True)
    wb.save(output_path)
    print(f"[EXCEL] Saved → {output_path}")

def _petro_source_text():
    parts = [SOURCE]
    if PETRO_LABEL:
        parts.append(f"label {PETRO_LABEL}")
    parts.append("từ khóa: " + ", ".join(PETRO_KEYWORDS))
    return " | ".join(parts)

def _export_petro(state):
    """Dựng lại file Excel Petro ĐẦY ĐỦ từ sổ theo dõi (gồm cả các lần chạy trước). Trả về số dòng."""
    _petro_dedupe(state, verbose=False)              # đảm bảo bản trùng đã được đánh dấu trước khi xuất
    rows = state.petro_rows(include_duplicates=False)
    if not rows:
        return 0
    def pdf_text(r):
        if not r["ma_tra_cuu"]:
            return "—"                       # không có mã tra cứu -> không thể tải
        if (r.get("pdf_status") or "PENDING") == "DONE":
            return r.get("pdf_file") or "Đã tải (thủ công)"
        return "Chưa tải"

    data = [{
        "date": r["header_date"] or r["mail_date"] or "",
        "sender": r["sender"] or "",
        "subject": r["subject"] or "",
        "so_hoa_don": r["so_hoa_don"] or "",
        "ma_tra_cuu": r["ma_tra_cuu"] or "",
        "pdf_file": pdf_text(r),
    } for r in rows]
    try:
        export_petro_excel(data, PETRO_OUTPUT, _petro_source_text(), dup_count=state.petro_duplicate_count())
    except PermissionError:
        print(f"[WARN] Không ghi được {PETRO_OUTPUT} (đang mở trong Excel?). Dữ liệu vẫn an toàn trong sổ theo dõi; "
              "đóng file rồi chạy lại với --export-petro để xuất lại.")
    except Exception as e:
        print(f"[WARN] Không xuất được Excel Petro | ERR={repr(e)}")
    return len(rows)


# ---------- Petro: tải PDF theo mã tra cứu (gộp từ down_petro.py) ----------
def _rel_to_download_dir(path):
    """Đường dẫn file tương đối so với thư mục tải chính (để --verify kiểm tra được)."""
    try:
        return os.path.relpath(path, DOWNLOAD_DIR)
    except ValueError:  # khác ổ đĩa trên Windows
        return path

def confirm_start(seconds):
    """
    Đếm ngược tự bắt đầu tải Petro. Enter = bắt đầu ngay, n = bỏ qua, không bấm gì = tự bắt đầu.
    Không dùng input() để không bị kẹt nếu cửa sổ dòng lệnh không nhận phím Enter.
    Trả về True nếu nên bắt đầu, False nếu người dùng chọn bỏ qua.
    """
    import time
    hint = "Enter = bắt đầu ngay, n = bỏ qua"

    try:
        import msvcrt  # Windows: đọc phím trực tiếp, không cần nhấn Enter sau chữ n
    except ImportError:
        msvcrt = None

    end = time.time() + seconds
    if msvcrt:
        last = None
        while time.time() < end:
            left = int(end - time.time()) + 1
            if left != last:
                print(f"\r[STEP 6] Tự động mở trình duyệt sau {left:>2}s ... ({hint})   ", end="", flush=True)
                last = left
            if msvcrt.kbhit():
                ch = msvcrt.getwch()
                if ch.lower() == "n":
                    print()
                    return False
                if ch in ("\r", "\n", " "):
                    print()
                    return True
            time.sleep(0.1)
        print()
        return True

    import select  # Linux / macOS
    while True:
        left = end - time.time()
        if left <= 0:
            print()
            return True
        print(f"\r[STEP 6] Tự động mở trình duyệt sau {int(left) + 1:>2}s ... ({hint}, rồi Enter)   ", end="", flush=True)
        ready, _, _ = select.select([sys.stdin], [], [], min(1.0, left))
        if ready:
            line = sys.stdin.readline()
            print()
            return not line.strip().lower().startswith("n")

# ---------- Petro: tự giải CAPTCHA (ddddocr), nhập tay dự phòng ----------
# Selector trang tra cứu hoadon.petrolimex.com.vn (lấy theo core_engine.py; nếu trang đổi giao diện thì sửa ở đây)
PETRO_SEL_CODE = "#strFkey"
PETRO_SEL_CAPTCHA_IMG = "#SearchformByfkey .captcha_img"
PETRO_SEL_CAPTCHA_INPUT = "#captch"
PETRO_SEL_SUBMIT = "#SearchformByfkey input[name='submit']"
PETRO_SEL_MESSAGE = "#messagewrapper"
PETRO_SEL_PDF = "a[onclick*='ajxCall4Portal']"      # bắt cả ajxCall4PortalPDF và ajxCall4Portal1PDF
PETRO_SEL_XML = "a[href*='DownloadXml']"
PETRO_AUTO_WAIT = 20                                # số giây chờ kết quả sau khi tự bấm Tìm kiếm
PETRO_AUTO_FAIL_STOP = 3                            # chế độ không người: dừng sau N mã liên tiếp tự giải CAPTCHA thất bại
_PETRO_OCR = None
_PETRO_OCR_TRIED = False

class _PetroBrowserClosed(Exception):
    pass

def _is_closed_error(e):
    msg = str(e)
    return any(k in msg for k in ("has been closed", "Target closed", "Browser closed", "Target page, context or browser"))

def _petro_get_ocr():
    """Khởi tạo ddddocr 1 lần. Trả về None nếu chưa cài / lỗi (khi đó dùng nhập tay)."""
    global _PETRO_OCR, _PETRO_OCR_TRIED
    if _PETRO_OCR_TRIED:
        return _PETRO_OCR
    _PETRO_OCR_TRIED = True
    try:
        import ddddocr
        try:
            _PETRO_OCR = ddddocr.DdddOcr(show_ad=False)
        except TypeError:                           # bản ddddocr cũ không có tham số show_ad
            _PETRO_OCR = ddddocr.DdddOcr()
        print("[STEP 6] Đã bật tự giải CAPTCHA (ddddocr).")
    except Exception as e:
        print(f"[WARN] Không dùng được ddddocr để tự giải CAPTCHA (cài bằng: pip install ddddocr) | ERR={repr(e)}")
        _PETRO_OCR = None
    return _PETRO_OCR

def _petro_classify_message(msg):
    """Phân loại thông báo của trang: bad_captcha | not_found | message | None (không có thông báo)."""
    m = _ascii_lower(msg)
    if not m:
        return None
    if "ma xac nhan" in m and ("khong" in m or "sai" in m):
        return "bad_captcha"
    if "khong ton tai" in m or "khong tim thay" in m:
        return "not_found"
    return "message"

async def _petro_page_message(page):
    try:
        el = page.locator(PETRO_SEL_MESSAGE).first
        if await el.count() and await el.is_visible():
            return (await el.inner_text()).strip()
    except Exception as e:
        if _is_closed_error(e):
            raise _PetroBrowserClosed() from e
    return ""

async def _petro_wait_result(page, timeout_s, auto):
    """
    Chờ kết quả tra cứu. Trả về (kết quả, thông báo):
      found        - đã có nút tải PDF
      not_found    - trang báo không tồn tại hóa đơn
      bad_captcha  - trang báo sai mã xác nhận (chỉ trả về ở chế độ tự động; nhập tay thì chờ người nhập lại)
      message      - thông báo khác của trang (chỉ ở chế độ tự động)
      timeout      - hết thời gian chờ
    """
    loop = asyncio.get_running_loop()
    end = loop.time() + timeout_s
    btn = page.locator(PETRO_SEL_PDF)
    last_msg = ""
    while loop.time() < end:
        try:
            if await btn.count() and await btn.last.is_visible():
                return "found", ""
        except Exception as e:                      # trang đang chuyển -> thử lại vòng sau
            if _is_closed_error(e):
                raise _PetroBrowserClosed() from e
        msg = await _petro_page_message(page)
        kind = _petro_classify_message(msg)
        if msg:
            last_msg = msg
        if kind == "not_found":
            return kind, msg
        if auto and kind in ("bad_captcha", "message"):
            return kind, msg
        await asyncio.sleep(0.4)
    return "timeout", last_msg

async def _petro_auto_attempt(page, ocr, code):
    """1 lần tự giải: mở trang, chụp CAPTCHA, nhận diện, điền và bấm Tìm kiếm. Trả về (kết quả, thông báo, chữ CAPTCHA)."""
    await page.goto(PETRO_LOOKUP_URL, wait_until="domcontentloaded", timeout=30000)
    img = page.locator(PETRO_SEL_CAPTCHA_IMG).first
    await img.wait_for(state="visible", timeout=15000)
    try:                                            # chờ ảnh tải xong (nếu là thẻ <img>)
        await img.evaluate("(el) => el.tagName !== 'IMG' || el.complete ? true : "
                           "new Promise(r => { el.onload = () => r(true); setTimeout(() => r(false), 5000); })")
    except Exception:
        pass
    cap_text = re.sub(r"\s+", "", ocr.classification(await img.screenshot()) or "")
    if not cap_text:
        return "bad_captcha", "OCR không đọc được CAPTCHA", cap_text
    await page.fill(PETRO_SEL_CODE, code)
    await page.fill(PETRO_SEL_CAPTCHA_INPUT, cap_text)
    await page.locator(PETRO_SEL_SUBMIT).first.click()
    kind, msg = await _petro_wait_result(page, PETRO_AUTO_WAIT, auto=True)
    return kind, msg, cap_text

async def _petro_lookup(page, code, ocr, manual_ok):
    """
    Tra cứu 1 mã. Tự giải CAPTCHA trước (thử cả mã có/không dấu *), không được thì chuyển sang nhập tay.
    Trả về (kết quả, chi tiết, mã đã dùng). Kết quả: found | not_found | message | captcha_failed | manual_timeout
    """
    variants = [code]
    alt = code.rstrip("*") if code.endswith("*") else code + "*"
    if alt and alt != code:
        variants.append(alt)

    if ocr:
        last, exhausted = None, False
        for v in variants:
            outcome = None
            for attempt in range(1, PETRO_CAPTCHA_ATTEMPTS + 1):
                kind, msg, cap = await _petro_auto_attempt(page, ocr, v)
                if kind == "found":
                    return "found", f"tự giải CAPTCHA (lần {attempt})", v
                if kind in ("not_found", "message"):
                    outcome = (kind, msg, v)
                    break
                why = "trang không phản hồi" if kind == "timeout" else f"CAPTCHA '{cap}' chưa đúng"
                print(f"   Lần {attempt}/{PETRO_CAPTCHA_ATTEMPTS}: {why}, thử lại...")
                await page.wait_for_timeout(600)
            if outcome is None:                     # hết lượt tự giải -> không thử biến thể nữa, chuyển nhập tay
                exhausted = True
                break
            last = outcome
            if v != variants[-1]:
                print(f"   Trang báo: {outcome[1] or outcome[0]} -> thử lại với mã {variants[-1]}")
        if not exhausted:
            return last
        if not manual_ok:
            return "captcha_failed", f"Tự giải CAPTCHA không thành công sau {PETRO_CAPTCHA_ATTEMPTS} lần", code
        print("   Tự giải CAPTCHA không thành công -> chuyển sang NHẬP TAY.")
    elif not manual_ok:
        return "captcha_failed", "Không có ddddocr và không có người nhập CAPTCHA", code

    # ----- Nhập tay -----
    await page.goto(PETRO_LOOKUP_URL, wait_until="domcontentloaded", timeout=30000)
    await page.fill(PETRO_SEL_CODE, code)
    try:
        await page.focus(PETRO_SEL_CAPTCHA_INPUT)
    except Exception:
        pass
    print(f">>> VUI LÒNG NHẬP CAPTCHA TRÊN TRÌNH DUYỆT VÀ BẤM TÌM KIẾM (chờ tối đa {PETRO_CAPTCHA_TIMEOUT}s)...")
    kind, msg = await _petro_wait_result(page, PETRO_CAPTCHA_TIMEOUT, auto=False)
    if kind == "found":
        return "found", "nhập tay", code
    if kind == "not_found":
        return kind, msg, code
    return "manual_timeout", f"Hết {PETRO_CAPTCHA_TIMEOUT}s chờ nhập CAPTCHA / không thấy nút tải PDF", code

async def _petro_row_meta(page):
    """Lấy Ký hiệu / Số HĐ / Công ty / Tổng tiền / Ngày xuất từ bảng kết quả (chỉ để đặt tên file & in log)."""
    meta = {}
    try:
        rows = page.locator("table.table tr")
        for i in range(await rows.count()):
            cells = rows.nth(i).locator("td")
            if await cells.count() >= 8:
                vals = [(await cells.nth(k).inner_text()).strip() for k in range(2, 8)]
                meta = dict(zip(("cong_ty", "mau_so", "ky_hieu", "so_hd", "tong_tien", "ngay_xuat"), vals))
                break
    except Exception:
        pass
    return meta

async def _petro_save_files(page, code):
    """Tải PDF (và XML nếu bật). Trả về (tên file PDF đã lưu, đường dẫn, tên file XML hoặc "", metadata)."""
    meta = await _petro_row_meta(page)
    btn = page.locator(PETRO_SEL_PDF).last
    async with page.expect_download(timeout=40000) as download_info:
        await btn.click(force=True)
    download = await download_info.value

    name = download.suggested_filename or ""
    if not name or name.lower() in ("hoadon.pdf", "download.pdf", "invoice.pdf"):
        so = meta.get("so_hd") or code.rstrip("*")
        name = f"HD_{meta.get('ky_hieu') or 'Petrolimex'}_{so}.pdf"
    name = re.sub(r'[\\/:*?"<>|]+', "_", name)
    file_name = get_unique_filename(PETRO_DOWNLOAD_DIR, name)
    save_path = os.path.join(PETRO_DOWNLOAD_DIR, file_name)
    await download.save_as(save_path)

    xml_name = ""
    if PETRO_XML:
        try:
            xml_btn = page.locator(PETRO_SEL_XML).last
            if await xml_btn.count():
                async with page.expect_download(timeout=20000) as xml_info:
                    await xml_btn.click(force=True)
                xml_dl = await xml_info.value
                xml_base = xml_dl.suggested_filename or (os.path.splitext(file_name)[0] + ".xml")
                xml_name = get_unique_filename(PETRO_DOWNLOAD_DIR, re.sub(r'[\\/:*?"<>|]+', "_", xml_base))
                await xml_dl.save_as(os.path.join(PETRO_DOWNLOAD_DIR, xml_name))
            else:
                print("   [WARN] Không thấy nút tải XML.")
        except Exception as e:
            if _is_closed_error(e):
                raise
            print(f"   [WARN] Tải XML lỗi (PDF vẫn đã lưu) | ERR={repr(e)[:200]}")
    return file_name, save_path, xml_name, meta

async def run_petro_downloads(state, prompt=True):
    """
    STEP 6 - tải PDF cho các hóa đơn Petro chưa tải. CHẠY CUỐI CÙNG.
      - Tự giải CAPTCHA bằng ddddocr (thử tối đa --petro-captcha-attempts lần, thử cả mã có/không dấu *).
      - Tự giải không được + có người ngồi máy -> chờ bạn nhập tay trên trình duyệt (như cách cũ).
      - Không có người ngồi máy (chạy lịch tự động) -> chạy ẩn, chỉ tự giải; mã không giải được để lần sau.
    """
    stats = {"done": 0, "failed": 0, "not_found": 0, "auto": 0, "manual": 0, "remaining": 0, "stopped": False}
    _petro_dedupe(state, verbose=not prompt)         # bỏ qua hóa đơn đã có trong file tổng / đã trùng email khác
    all_pending = state.petro_pending_downloads()
    stats["remaining"] = len(all_pending)
    if not all_pending:
        print("[STEP 6 DONE] Không có hóa đơn Petro nào cần tải PDF.")
        return stats

    batch = all_pending[:PETRO_LIMIT] if PETRO_LIMIT else all_pending
    print(f"[STEP 6] Petro: {len(all_pending)} hóa đơn chưa tải PDF"
          + (f" (lần này tải {len(batch)}, giới hạn --petro-limit {PETRO_LIMIT})" if len(batch) < len(all_pending) else ""))
    print(f"[STEP 6] Thư mục lưu: {PETRO_DOWNLOAD_DIR}")

    hint_cmd = "python main.py --download-petro" + (f" --profile {ARGS_PROFILE}" if ARGS_PROFILE else "")
    interactive = bool(sys.stdin) and sys.stdin.isatty()
    ocr = _petro_get_ocr() if PETRO_AUTO_CAPTCHA else None
    manual_ok = interactive and not PETRO_HEADLESS
    if not ocr and not manual_ok:
        print("[STEP 6] Không tự giải được CAPTCHA (thiếu ddddocr hoặc đã tắt) và không có người nhập tay -> bỏ qua. "
              f"Khi có mặt, chạy: {hint_cmd}")
        return stats

    if ocr and manual_ok:
        print(f"[STEP 6] Chế độ: TỰ GIẢI CAPTCHA (tối đa {PETRO_CAPTCHA_ATTEMPTS} lần/mã), không được thì chờ bạn NHẬP TAY "
              f"(tối đa {PETRO_CAPTCHA_TIMEOUT}s/mã).")
    elif ocr:
        print(f"[STEP 6] Chế độ: CHỈ TỰ GIẢI CAPTCHA, chạy ẩn (tối đa {PETRO_CAPTCHA_ATTEMPTS} lần/mã). "
              "Mã không giải được sẽ để lại cho lần sau.")
    else:
        print("[STEP 6] Chế độ: NHẬP TAY - trình duyệt sẽ mở, bạn cần nhập CAPTCHA và bấm Tìm kiếm cho từng mã "
              f"(chờ tối đa {PETRO_CAPTCHA_TIMEOUT}s/mã).")
        if prompt and not confirm_start(PETRO_START_DELAY):
            print(f"[STEP 6] Đã bỏ qua tải PDF Petro (các hóa đơn vẫn ở trạng thái chưa tải). Khi sẵn sàng, chạy: {hint_cmd}")
            return stats

    ensure_dir(PETRO_DOWNLOAD_DIR)
    headless = not manual_ok                          # cần nhập tay -> phải thấy trình duyệt
    print(f"[STEP 6] Đang mở trình duyệt Chromium{' (chạy ẩn)' if headless else ''}...")
    manual_timeouts_in_row = 0
    auto_fail_in_row = 0

    async with async_playwright() as p:
        try:
            browser = await p.chromium.launch(headless=headless)
        except Exception as e:
            print(f"[ERROR] Không mở được trình duyệt | ERR={repr(e)}")
            return stats
        context = await browser.new_context(accept_downloads=True)
        page = await context.new_page()
        if manual_ok:
            print("[STEP 6] Trình duyệt đã mở - nếu cần nhập tay, hãy chuyển sang cửa sổ Chromium.")

        try:
            for idx, row in enumerate(batch, 1):
                mid, code = row["message_id"], row["ma_tra_cuu"]
                print(f"\n[{idx}/{len(batch)}] Đang xử lý mã: {code}")
                try:
                    kind, detail, used_code = await _petro_lookup(page, code, ocr, manual_ok)

                    if kind == "found":
                        print(f"   Đã tìm thấy hóa đơn ({detail}"
                              + (f", mã dùng: {used_code}" if used_code != code else "") + "). Đang lưu file...")
                        try:
                            file_name, save_path, xml_name, meta = await _petro_save_files(page, used_code)
                        except PlaywrightTimeoutError:
                            stats["failed"] += 1
                            reason = "Đã có nút tải nhưng không nhận được file (hết thời gian chờ tải)"
                            print(f"   -> LỖI cho mã {code}: {reason}")
                            state.set_petro_pdf(mid, FAILED, error=reason)
                            continue
                        state.set_petro_pdf(mid, DONE, pdf_file=_rel_to_download_dir(save_path))
                        stats["done"] += 1
                        stats["manual" if detail == "nhập tay" else "auto"] += 1
                        manual_timeouts_in_row = auto_fail_in_row = 0
                        info = " | ".join(f"{k}={v}" for k, v in (
                            ("Ký hiệu", meta.get("ky_hieu")), ("Số HĐ", meta.get("so_hd")),
                            ("Tổng tiền", meta.get("tong_tien")), ("Ngày", meta.get("ngay_xuat"))) if v)
                        print(f"   -> ĐÃ TẢI XONG: {file_name}" + (f" + {xml_name}" if xml_name else "")
                              + (f" ({info})" if info else ""))

                    elif kind == "not_found":
                        stats["failed"] += 1
                        stats["not_found"] += 1
                        manual_timeouts_in_row = auto_fail_in_row = 0
                        reason = f"Không tồn tại hóa đơn với mã này (trang báo: {detail})"[:300]
                        print(f"   -> {reason}")
                        state.set_petro_pdf(mid, FAILED, error=reason, not_found=True)
                        if state.petro_not_found_count(mid) >= PETRO_NOT_FOUND_MAX:
                            print(f"   -> Đã {PETRO_NOT_FOUND_MAX} lần báo không tồn tại -> KHÔNG tự thử lại nữa. "
                                  "Kiểm tra lại mã trong email; muốn thử lại thì 'Đánh dấu xử lý lại' email này.")

                    elif kind == "message":
                        stats["failed"] += 1
                        manual_timeouts_in_row = auto_fail_in_row = 0
                        reason = f"Trang tra cứu báo: {detail}"[:300]
                        print(f"   -> LỖI cho mã {code}: {reason}")
                        state.set_petro_pdf(mid, FAILED, error=reason)

                    elif kind == "manual_timeout":
                        stats["failed"] += 1
                        manual_timeouts_in_row += 1
                        print(f"   -> LỖI cho mã {code}: {detail}")
                        state.set_petro_pdf(mid, FAILED, error=detail)
                        if manual_timeouts_in_row >= 2:
                            print("\n[STEP 6] 2 mã liên tiếp không có thao tác nhập CAPTCHA -> dừng để tránh chờ vô ích. "
                                  "Các mã còn lại giữ nguyên trạng thái chưa tải.")
                            stats["stopped"] = True
                            break

                    else:  # captcha_failed (chế độ không người)
                        stats["failed"] += 1
                        auto_fail_in_row += 1
                        print(f"   -> LỖI cho mã {code}: {detail}")
                        state.set_petro_pdf(mid, FAILED, error=detail)
                        if auto_fail_in_row >= PETRO_AUTO_FAIL_STOP:
                            print(f"\n[STEP 6] {PETRO_AUTO_FAIL_STOP} mã liên tiếp tự giải CAPTCHA thất bại -> dừng "
                                  f"(có thể trang đổi CAPTCHA/chặn truy cập). Khi có mặt, chạy: {hint_cmd}")
                            stats["stopped"] = True
                            break

                except _PetroBrowserClosed:
                    print("\n[STEP 6] Trình duyệt đã bị đóng -> dừng tải. Các mã còn lại giữ nguyên trạng thái chưa tải.")
                    stats["stopped"] = True
                    break
                except Exception as e:
                    if _is_closed_error(e):
                        print("\n[STEP 6] Trình duyệt đã bị đóng -> dừng tải. Các mã còn lại giữ nguyên trạng thái chưa tải.")
                        stats["stopped"] = True
                        break
                    stats["failed"] += 1
                    msg = str(e)
                    print(f"   -> LỖI cho mã {code}.")
                    print(f"   -> Chi tiết lỗi: {msg[:300]}")
                    state.set_petro_pdf(mid, FAILED, error=msg[:300])

                await asyncio.sleep(0.5)             # nghỉ ngắn giữa các mã để không dồn dập vào máy chủ
        finally:
            try:
                await browser.close()
            except Exception:
                pass

    stats["remaining"] = len(state.petro_pending_downloads())
    print(f"\n[STEP 6 DONE] Petro PDF - Đã tải: {stats['done']} (tự giải: {stats['auto']}, nhập tay: {stats['manual']}) | "
          f"Lỗi: {stats['failed']} (không tồn tại: {stats['not_found']}) | Còn chưa tải: {stats['remaining']}")
    return stats

# ================== EXCEL LỖI / BỎ QUA / CHƯA XỬ LÝ (để làm thủ công) ==================
_REASON_HINTS = [
    ("PDF_ATTACHMENT_NOT_DOWNLOADED", "Có PDF đính kèm nhưng tải lỗi -> mở email và tải tay"),
    ("SUPPORTED_SUPPLIER_BUT_NO_LINK", "Nhận diện được nhà cung cấp nhưng không tìm thấy link tải -> mở email lấy hóa đơn"),
    ("UNSUPPORTED_SUPPLIER", "Nhà cung cấp chưa được hỗ trợ -> xử lý tay"),
    ("NO_PDF_NO_SUPPLIER_LINK", "Không có PDF/link hóa đơn (có thể không phải mail hóa đơn) -> kiểm tra email"),
    ("PETRO_KEYWORD_BUT_NO_INVOICE_FIELDS", "Khớp từ khóa Petro nhưng không có số hóa đơn/mã tra cứu -> kiểm tra email"),
    ("PETRO_MISSING_DATA", "Petrolimex thiếu số hóa đơn/mã tra cứu -> đọc email rồi nhập tay"),
    ("PETRO_FETCH_ERROR", "Không đọc được email Petro -> mở email kiểm tra hoặc chạy lại"),
    ("SCAN_READ_ERROR", "Đọc email lỗi (mạng/quota) -> chạy lại chương trình"),
]

def _reason_hint(reason):
    reason = reason or ""
    for prefix, hint in _REASON_HINTS:
        if reason.startswith(prefix):
            return hint
    return "Tải/xử lý lỗi -> mở link hoặc email để làm tay" if reason else ""

def _export_issues(state, issues, url_by_id=None, xml_rows=None):
    """
    Xuất 1 file Excel gồm các mục cần làm THỦ CÔNG:
      - email LỖI / BỊ BỎ QUA trong lần chạy này (issues)
      - hóa đơn Petro đã có mã tra cứu nhưng CHƯA TẢI được PDF
    Không có mục nào -> không tạo file. Trả về số dòng đã xuất.
    """
    url_by_id = url_by_id or {}
    petro_by_id, petro_pending = {}, []
    try:
        for r in state.petro_rows(include_duplicates=False):
            mid = r.get("message_id")
            if mid:
                petro_by_id[mid] = r
            if r.get("ma_tra_cuu") and (r.get("pdf_status") or "PENDING") != "DONE":
                petro_pending.append(r)
    except Exception as e:
        print(f"[WARN] Không đọc được danh sách Petro chưa tải | ERR={repr(e)}")

    rows = []
    kind_label = {FAILED: "LỖI", SKIPPED: "BỎ QUA"}
    for eid, it in issues.items():
        real_id = "" if str(eid).startswith("UNKNOWN") else eid
        p = petro_by_id.get(real_id) or {}
        rows.append({
            "kind": kind_label.get(it["status"], it["status"]),
            "supplier": it.get("supplier") or "",
            "date": it.get("mail_date") or "",
            "subject": it.get("subject") or "",
            "so_hoa_don": p.get("so_hoa_don") or "",
            "ma_tra_cuu": p.get("ma_tra_cuu") or "",
            "reason": str(it.get("reason") or "")[:500],
            "hint": _reason_hint(it.get("reason")),
            "url": url_by_id.get(real_id, ""),
            "email_id": real_id,
        })

    for r in petro_pending:
        mid = r.get("message_id") or ""
        if mid in issues:
            continue                                   # đã có dòng lỗi của email này rồi
        pending = (r.get("pdf_status") or "PENDING") == "PENDING"
        rows.append({
            "kind": "PETRO CHƯA TẢI PDF",
            "supplier": "petrolimex",
            "date": r.get("header_date") or r.get("mail_date") or "",
            "subject": r.get("subject") or "",
            "so_hoa_don": r.get("so_hoa_don") or "",
            "ma_tra_cuu": r.get("ma_tra_cuu") or "",
            "reason": ("Đã có mã tra cứu nhưng chưa tải PDF" if pending else
                       f"Trang tra cứu báo không tồn tại hóa đơn {PETRO_NOT_FOUND_MAX} lần -> đã thôi tự thử lại, kiểm tra mã"
                       if (r.get("pdf_not_found") or 0) >= PETRO_NOT_FOUND_MAX else "Tải PDF lỗi ở lần trước"),
            "hint": f"Tra cứu bằng mã tra cứu tại {PETRO_LOOKUP_URL} rồi tải PDF (hoặc chạy: python main.py --download-petro)",
            "url": PETRO_LOOKUP_URL,
            "email_id": mid,
        })

    xml_rows = xml_rows or []
    if not rows and not xml_rows:
        print("[ERRORS] Không có email lỗi / bỏ qua / chưa xử lý -> không tạo file Excel lỗi.")
        return 0

    order = {"LỖI": 0, "PETRO CHƯA TẢI PDF": 1, "BỎ QUA": 2}
    rows.sort(key=lambda x: (order.get(x["kind"], 9), x["supplier"], x["date"]))

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cần xử lý thủ công"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headers = ["STT", "Loại", "Nhà cung cấp", "Ngày mail", "Tiêu đề email", "Số hóa đơn", "Mã tra cứu",
               "Lý do / chi tiết lỗi", "Gợi ý xử lý", "Link tải / tra cứu", "Mở email", "Email ID"]
    widths = [6, 20, 18, 24, 45, 16, 18, 45, 50, 40, 14, 20]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", start_color="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 28

    kind_fill = {"LỖI": "FFC7CE", "BỎ QUA": "FFEB9C", "PETRO CHƯA TẢI PDF": "DDEBF7"}
    for idx, r in enumerate(rows, 1):
        vals = [idx, r["kind"], r["supplier"], r["date"], r["subject"], r["so_hoa_don"], r["ma_tra_cuu"],
                r["reason"], r["hint"], r["url"], "Mở email" if r["email_id"] else "", r["email_id"]]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=idx + 1, column=col, value=v)
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(horizontal="center" if col == 1 else "left", vertical="top", wrap_text=col in (5, 8, 9, 10))
            c.border = border
            if col in (6, 7):
                c.number_format = "@"                  # giữ số 0 đầu của số hóa đơn / mã tra cứu
        ws.cell(row=idx + 1, column=2).fill = PatternFill("solid", start_color=kind_fill.get(r["kind"], "FFFFFF"))
        ws.cell(row=idx + 1, column=2).font = Font(name="Arial", size=10, bold=True)
        if r["email_id"]:
            link_cell = ws.cell(row=idx + 1, column=11)
            link_cell.hyperlink = gmail_link(r["email_id"])
            link_cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(headers)).column_letter}{len(rows) + 1}"

    if xml_rows:                                       # --keep-xml: sheet riêng liệt kê email thiếu XML / lỗi XML
        wx = wb.create_sheet("Thiếu XML")
        xh = ["STT", "Tình trạng", "Loại email", "Nhà cung cấp", "Ngày mail", "Tiêu đề email", "File PDF",
              "Chi tiết lỗi", "Gợi ý xử lý", "Mở email", "Email ID"]
        xw = [6, 16, 22, 18, 20, 45, 36, 40, 50, 14, 20]
        for i, (h, w) in enumerate(zip(xh, xw), 1):
            c = wx.cell(row=1, column=i, value=h)
            c.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
            c.fill = PatternFill("solid", start_color="1F4E79")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border
            wx.column_dimensions[c.column_letter].width = w
        for idx, r in enumerate(xml_rows, 1):
            vals = [idx, r["status"], r["kind"], r["supplier"], r["date"], r["subject"], r["pdf"], r["error"],
                    r["hint"], "Mở email" if r["email_id"] else "", r["email_id"]]
            for col, v in enumerate(vals, 1):
                c = wx.cell(row=idx + 1, column=col, value=v)
                c.font = Font(name="Arial", size=10)
                c.alignment = Alignment(horizontal="center" if col == 1 else "left", vertical="top", wrap_text=col in (6, 8, 9))
                c.border = border
            wx.cell(row=idx + 1, column=2).fill = PatternFill("solid", start_color="FFC7CE" if r["status"] == "LỖI XML" else "FFEB9C")
            wx.cell(row=idx + 1, column=2).font = Font(name="Arial", size=10, bold=True)
            if r["email_id"]:
                link_cell = wx.cell(row=idx + 1, column=10)
                link_cell.hyperlink = gmail_link(r["email_id"])
                link_cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")
        wx.freeze_panes = "A2"
        wx.auto_filter.ref = f"A1:{wx.cell(row=1, column=len(xh)).column_letter}{len(xml_rows) + 1}"

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(ERROR_DIR, f"loi_can_xu_ly_{ARGS_PROFILE or 'default'}_{stamp}.xlsx")
    try:
        ensure_dir(ERROR_DIR)
        wb.save(path)
    except Exception as e:
        print(f"[WARN] Không xuất được Excel lỗi tại {path} | ERR={repr(e)}")
        return 0
    by_kind = {}
    for r in rows:
        by_kind[r["kind"]] = by_kind.get(r["kind"], 0) + 1
    print(f"[ERRORS] Có {len(rows)} mục cần xử lý thủ công" + (f" ({' | '.join(f'{k}: {v}' for k, v in by_kind.items())})" if by_kind else ""))
    if xml_rows:
        print(f"[ERRORS] Sheet 'Thiếu XML': {len(xml_rows)} email không có XML / lỗi XML")
    print(f"[ERRORS] Đã xuất file Excel: {path}")
    return len(rows)

# ================== SỔ THEO DÕI / KIỂM TRA THỦ CÔNG ==================
def _open_state():
    return EmailState(STATE_PATH, max_attempts=MAX_ATTEMPTS)

def _csv_path():
    return os.path.splitext(STATE_PATH)[0] + ".csv"

def _export_csv(state):
    path = _csv_path()
    try:
        n = state.export_csv(path)
        print(f"[STATE] Đã xuất {n} dòng ra CSV: {path}")
    except PermissionError:
        print(f"[WARN] Không ghi được {path} (đang mở trong Excel?). Đóng file rồi chạy lại --report để xuất lại.")
    except Exception as e:
        print(f"[WARN] Không xuất được CSV | ERR={repr(e)}")

def build_date_query():
    """Điều kiện khoảng ngày (--since / --until) theo cú pháp Gmail; rỗng nếu không giới hạn."""
    parts = []
    if SINCE_TS:
        parts.append(f"after:{SINCE_TS}")
    if UNTIL_TS:
        parts.append(f"before:{UNTIL_TS}")
    return " ".join(parts)

def build_gmail_query():
    """Trả về Gmail search query, hoặc None nếu quét theo label."""
    if SOURCE == "label":
        return None
    parts = []
    if SOURCE == "inbox":
        parts.append("in:inbox")
    else:  # all
        parts.append("-in:sent -in:drafts")
    if EXTRA_QUERY:
        parts.append(EXTRA_QUERY)
    date_q = build_date_query()
    if date_q:
        parts.append(date_q)
    return " ".join(parts)

def _print_rows(state, rows):
    for r in rows:
        files = state.files_of(r)
        exhausted = ""
        if r["status"] == FAILED and r["attempts"] >= state.max_attempts:
            exhausted = " (HẾT LƯỢT THỬ)"
        print(
            f"  - EMAIL={r['message_id']} | SUPPLIER={r['supplier'] or '-'} | "
            f"DATE={r['mail_date'] or '-'} | ATTEMPTS={r['attempts']}/{state.max_attempts}{exhausted}"
        )
        print(f"    SUBJECT={r['subject'] or '-'}")
        if r["reason"]:
            print(f"    REASON={r['reason']}")
        if files:
            print(f"    FILES={'; '.join(files)}")
        print(f"    OPEN={gmail_link(r['message_id'])}")

def print_report(state, status=None, limit=50):
    counts = state.counts()
    total = sum(counts.values())
    print(f"[REPORT] DB: {STATE_PATH}")
    print(f"[REPORT] Tổng: {total} | " + " | ".join(f"{s}: {counts.get(s, 0)}" for s in ALL_STATUSES))
    if state.petro_count():
        print(f"[REPORT] Petro: {state.petro_count()} hóa đơn trong sổ | Excel: {PETRO_OUTPUT}")
        pc_ = state.petro_pdf_counts()
        print(f"[REPORT] Petro PDF: đã tải {pc_.get('DONE', 0)} | chưa tải {pc_.get('PENDING', 0)} | "
              f"lỗi lần trước {pc_.get('FAILED', 0)} | trùng (không tải) {pc_.get('DUPLICATE', 0)} | "
              f"không có mã tra cứu {pc_.get('NO_CODE', 0)}")
        if PETRO_MASTER:
            print(f"[REPORT] File tổng Petro: {PETRO_MASTER}")

    lim = limit if limit and limit > 0 else None
    statuses = [status.upper()] if status else [FAILED, PENDING]
    for st in statuses:
        n = counts.get(st, 0)
        if not n:
            continue
        note = f" (hiển thị {lim} mới nhất, dùng --limit 0 để xem hết)" if lim and n > lim else ""
        print(f"\n[{st}] {n} email{note}")
        _print_rows(state, state.rows(st, lim))

    if not status:
        breakdown = state.reason_breakdown(SKIPPED)
        if breakdown:
            print("\n[SKIPPED] theo lý do (xem chi tiết: --report --status skipped):")
            for reason, n in breakdown:
                print(f"  - {reason or '(không rõ)'}: {n}")

def run_management_commands(args):
    """Các lệnh kiểm tra/chỉnh sổ thủ công - không cần quét Gmail. Trả về True nếu đã xử lý."""
    if not (args.mark_done or args.mark_pending or args.verify or args.report or args.export_petro
            or args.mark_petro_downloaded is not None):
        return False

    state = _open_state()
    try:
        for eid in (args.mark_done or []):
            existed = state.mark(eid, DONE)
            print(f"[MARK] {eid} -> DONE" + ("" if existed else " (email chưa có trong sổ, đã thêm mới)"))
        for eid in (args.mark_pending or []):
            existed = state.mark(eid, PENDING)
            print(f"[MARK] {eid} -> PENDING (sẽ được xử lý lại ở lần chạy sau)" + ("" if existed else " (email chưa có trong sổ, đã thêm mới)"))

        if args.verify:
            checked, no_info, problems = 0, 0, []       # tìm file cả trong thư mục con Cac_hang_khac / Khong_co_XML
            for row in state.rows(DONE):
                files = state.files_of(row)
                if not files:
                    no_info += 1
                    continue
                checked += 1
                missing = [f for f in files if not _find_pdf(f)]
                if missing:
                    row["missing"] = missing
                    problems.append(row)
            print(f"[VERIFY] Thư mục: {DOWNLOAD_DIR}")
            print(f"[VERIFY] Đã kiểm tra {checked} email DONE có ghi tên file | {no_info} email DONE không có thông tin file (seed/thủ công) | THIẾU FILE: {len(problems)}")
            for r in problems:
                print(f"  - EMAIL={r['message_id']} | THIẾU={'; '.join(r['missing'])}")
                print(f"    SUBJECT={r['subject'] or '-'}")
                print(f"    OPEN={gmail_link(r['message_id'])}")
                if args.fix:
                    state.mark(r["message_id"], PENDING)
            if problems:
                if args.fix:
                    print(f"[VERIFY] Đã đánh dấu PENDING {len(problems)} email. Chạy lại chương trình để tải lại.")
                else:
                    print("[VERIFY] Thêm --fix để tự đánh dấu PENDING các email này, hoặc dùng --mark-pending <ID>.")

        if args.mark_petro_downloaded is not None:
            keys = args.mark_petro_downloaded
            n = state.mark_petro_downloaded(keys)
            target = f"{len(keys)} mã/ID đã chọn" if keys else "TẤT CẢ hóa đơn đang chờ tải"
            print(f"[MARK] Petro: đã đánh dấu {n} hóa đơn là ĐÃ TẢI PDF ({target}) - sẽ không tải lại.")
            _export_petro(state)
            _petro_sync_master(state)

        if args.export_petro:
            n = _export_petro(state)
            _petro_sync_master(state)
            if not n:
                print("[PETRO] Sổ Petro đang trống, chưa có gì để xuất.")

        if args.report or args.mark_done or args.mark_pending or args.verify:
            if args.report:
                print()
                print_report(state, args.status, args.limit)
            _export_csv(state)
    finally:
        state.close()
    return True

def seed_from_label(service, state, label_name):
    label_id = get_label_id(service, label_name)
    if not label_id:
        print(f"[SEED] Không tìm thấy label '{label_name}' - bỏ qua bước seed.")
        return
    msgs = list_all_messages(service, label_id=label_id)
    added = state.seed([m["id"] for m in msgs], reason=f"SEEDED_FROM_LABEL:{label_name}")
    print(f"[SEED] Label '{label_name}' có {len(msgs)} email -> thêm {added} email vào sổ (DONE), {len(msgs) - added} email đã có sẵn trong sổ.")


# ================== MAIN ==================
async def _pipeline(args, creds, gmail_service, state):
    query = build_gmail_query()
    if query is None:
        source_text = f"Label: {LABEL_NAME}"
    else:
        source_text = f"Source: {SOURCE} | Query: {query}"
    print(f"[CONFIG] {source_text} | Download: {DOWNLOAD_DIR} | Token: {TOKEN_PATH} | State: {STATE_PATH}")
    if SINCE_TS or UNTIL_TS:
        d_from = datetime.fromtimestamp(SINCE_TS).strftime("%d/%m/%Y") if SINCE_TS else "từ đầu"
        d_to = (datetime.fromtimestamp(UNTIL_TS) - timedelta(days=1)).strftime("%d/%m/%Y") if UNTIL_TS else "nay"
        print(f"[CONFIG] Khoảng ngày quét: {d_from} → {d_to} (gồm cả ngày đầu và ngày cuối)")
    if KEEP_XML:
        print(f"[CONFIG] Lưu thêm file XML hóa đơn (đính kèm / trong ZIP) vào: {_xml_dir()}")

    run_counts = {DONE: 0, FAILED: 0, SKIPPED: 0}

    issues = {}            # email LỖI/BỊ BỎ QUA trong lần chạy này -> xuất Excel để làm thủ công
    link_url_by_id = {}    # email_id -> link tải (để ghi vào Excel lỗi)

    def record(eid, status, **kwargs):
        if eid and eid != "UNKNOWN":
            state.record(eid, status, **kwargs)
        run_counts[status] = run_counts.get(status, 0) + 1
        key = eid if eid and eid != "UNKNOWN" else f"UNKNOWN#{len(issues)}"
        if status in (FAILED, SKIPPED):
            issues[key] = {
                "status": status,
                "supplier": kwargs.get("supplier") or "",
                "reason": kwargs.get("reason") or "",
                "subject": kwargs.get("subject") or "",
                "mail_date": kwargs.get("mail_date") or "",
            }
        else:
            issues.pop(key, None)     # xử lý được ở lần thử sau trong cùng lần chạy -> không còn là lỗi

    # ---------- STEP 0: đọc sổ theo dõi ----------
    if args.seed_label:
        seed_from_label(gmail_service, state, args.seed_label)

    if args.force:
        skip_ids, include_ids = set(), set()
        print("[STATE] --force: bỏ qua sổ theo dõi, quét lại tất cả (có thể tạo file trùng).")
    else:
        skip_ids, include_ids = state.partition_ids(
            retry_failed=args.retry_failed, reprocess_skipped=args.reprocess_skipped
        )
        counts = state.counts()
        print("[STATE] Sổ hiện có: " + " | ".join(f"{s}: {counts.get(s, 0)}" for s in ALL_STATUSES))
        print(f"[STATE] Bỏ qua {len(skip_ids)} email đã xử lý | thử lại {len(include_ids)} email (FAILED/PENDING).")

    # Mail thuộc label Petro (kể cả đã rời Inbox) -> đưa vào danh sách quét nếu chưa xử lý xong
    petro_label_ids = set()
    if PETRO_LABEL:
        try:
            pl_id = get_label_id(gmail_service, PETRO_LABEL)
            if pl_id:
                petro_label_ids = {m["id"] for m in list_all_messages(gmail_service, label_id=pl_id, query=build_date_query() or None)}
                extra = petro_label_ids - skip_ids
                include_ids = set(include_ids) | extra
                print(f"[PETRO] Label '{PETRO_LABEL}': {len(petro_label_ids)} email ({len(extra)} email chưa xử lý).")
            else:
                print(f"[PETRO] Không có label '{PETRO_LABEL}' - chỉ nhận diện Petro theo từ khóa: {', '.join(PETRO_KEYWORDS)}.")
        except Exception as e:
            print(f"[WARN] Không đọc được label Petro '{PETRO_LABEL}' | ERR={repr(e)}")

    print("\n[STEP 1] Đang quét email (Tích hợp thông tin PDF đính kèm)...")
    t0 = time.perf_counter()
    try:
        result = run_scan(
            label_name=LABEL_NAME,
            token_path=TOKEN_PATH,
            credentials_path=CREDENTIALS_PATH,
            query=query,
            skip_ids=skip_ids,
            include_ids=include_ids,
            keywords=PETRO_KEYWORDS,
            extra_query=build_date_query() or None,
        )
    except Exception as e:
        # Không quét được Gmail (mất mạng, token hết hạn, vượt giới hạn...) -> DỪNG và báo lỗi.
        # Trước đây chạy tiếp với danh sách rỗng nên giao diện vẫn báo "Đã chạy xong".
        raise ScanError(f"Không quét được Gmail: {e!r}") from e

    all_data = result.get("data", [])
    read_errors = result.get("read_errors", [])
    print(
        f"[STEP 1 DONE] Gmail liệt kê: {result.get('listed', len(all_data))} email | "
        f"bỏ qua vì đã xử lý: {result.get('skipped_known', 0)} | "
        f"cần xử lý lần này: {len(all_data) + len(read_errors)}"
    )
    print(f"[TIME] STEP 1 (quét email): {time.perf_counter() - t0:.1f}s")

    # Email đọc lỗi (mạng, quota...) -> ghi FAILED để lần sau thử lại
    for err in read_errors:
        record(err.get("id"), FAILED, reason=f"SCAN_READ_ERROR: {str(err.get('error'))[:200]}", commit=False)
    if read_errors:
        state.commit()
        print(f"[STEP 1] {len(read_errors)} email đọc lỗi -> đã ghi FAILED, sẽ thử lại lần chạy sau.")

    meta_by_id = {item["id"]: item for item in all_data}
    pdf_emails_ids = [item["id"] for item in all_data if item.get("has_pdf")]
    # File đính kèm đã lấy sẵn lúc quét -> STEP 2 và bước lấy XML không phải đọc lại email
    parts_by_id = {item["id"]: item["attachments"] for item in all_data if "attachments" in item}

    print(f"\n[STEP 2] Phát hiện {len(pdf_emails_ids)} email có file PDF đính kèm.")
    pdf_processed_ids = set()
    if pdf_emails_ids:
        print("[STEP 2] Đang tải đa luồng các file PDF trực tiếp...")
        t0 = time.perf_counter()
        downloaded, pdf_processed_ids = download_pdfs_parallel(pdf_emails_ids, creds, parts_by_id)
        print(f"[STEP 2 DONE] Đã tải xuống {downloaded} file đính kèm trực tiếp.")
        print(f"[TIME] STEP 2 (tải PDF đính kèm): {time.perf_counter() - t0:.1f}s")
        for eid in pdf_processed_ids:
            files = pop_tracked_files(eid)
            meta = meta_by_id.get(eid, {})
            record(eid, DONE, supplier="pdf_attachment", saved_count=len(files), files=files,
                   subject=meta.get("subject"), mail_date=meta.get("mail_date"), commit=False)
        state.commit()
    else:
        print("[STEP 2 DONE] Không có PDF đính kèm trực tiếp nào cần tải.")

    supplier_count = {"smartsign":0,"pvoil":0,"meinvoice":0,"easyinvoice":0,"fast":0,"ehoadon":0,"vnpt":0}
    link_items = []
    skipped_items = []
    petro_items = []   # xử lý SAU CÙNG (STEP 5)

    for item in all_data:
        raw_supplier = item.get("supplier")
        supplier = raw_supplier if raw_supplier else ""
        content = item.get("content", "")
        eid = item.get("id") or item.get("email_id") or "UNKNOWN"

        if eid in pdf_processed_ids:
            continue

        # Tiêu đề đã được lấy sẵn khi quét; chỉ gọi API riêng nếu thiếu
        subject = item.get("subject")
        if not subject:
            subject = "UNKNOWN_EMAIL_ID" if eid == "UNKNOWN" else get_email_subject(gmail_service, eid)
        mail_date = item.get("mail_date")

        link = None
        key = None

        if "easyinvoice" in supplier:
            link = extract_easyinvoice_link(content)
            key = "easyinvoice"
        elif "meinvoice.vn" in supplier:
            link = extract_meinvoice_link(content)
            key = "meinvoice"
        elif "pvoil" in supplier:
            link = extract_pvoil_link(content)
            key = "pvoil"
        elif "smartsign.com.vn" in supplier:
            link = extract_smartsign_link(content)
            key = "smartsign"
        elif "fast" in supplier or "fastsoftware" in supplier:
            link = extract_fast_link(content)
            key = "fast"
        elif "ehoadon.vn" in supplier:
            link = extract_ehoadon_link(content)
            key = "ehoadon"
        elif "vnpt-invoice.com.vn" in supplier or "vnpt" in supplier:
            link = extract_vnpt_link(content)
            key = "vnpt"

        if link:
            link_items.append({
                "supplier": key,
                "url": link,
                "email_id": eid,
                "subject": subject,
                "mail_date": mail_date,
                "content": content
            })
            supplier_count[key] += 1
        else:
            if item.get("has_pdf"):
                reason = "PDF_ATTACHMENT_NOT_DOWNLOADED"
            elif key is None and (eid in petro_label_ids or item.get("keyword_hit")):
                # Không thuộc nhà cung cấp đã hỗ trợ + thuộc label/từ khóa Petro -> để STEP 5 xử lý
                petro_items.append(item)
                continue
            elif supplier:
                reason = "SUPPORTED_SUPPLIER_BUT_NO_LINK" if key else "UNSUPPORTED_SUPPLIER"
            else:
                reason = "NO_PDF_NO_SUPPLIER_LINK"
            skipped_items.append({
                "email_id": eid,
                "subject": subject,
                "supplier": supplier or "UNKNOWN",
                "reason": reason,
            })
            # PDF đính kèm tải lỗi -> FAILED (sẽ thử lại). Các trường hợp còn lại -> SKIPPED (không quét lại vô ích).
            status = FAILED if reason == "PDF_ATTACHMENT_NOT_DOWNLOADED" else SKIPPED
            record(eid, status, supplier=supplier or None, reason=reason,
                   files=pop_tracked_files(eid), subject=subject, mail_date=mail_date, commit=False)
    state.commit()

    total_links = len(link_items)
    link_url_by_id.update({i["email_id"]: i["url"] for i in link_items})
    if KEEP_XML:
        # Email tải PDF qua link, hoặc không có PDF (có thể chỉ có XML/ZIP đính kèm) -> vẫn lấy XML đính kèm nếu có
        xml_check = [(i["email_id"], f"PDF qua link {i['supplier']}") for i in link_items if i["email_id"] != "UNKNOWN"]
        xml_check += [(i["email_id"], "Không có PDF") for i in skipped_items
                      if i["email_id"] != "UNKNOWN" and i.get("reason") != "PDF_ATTACHMENT_NOT_DOWNLOADED"]
        t0 = time.perf_counter()
        check_xml_for_emails(xml_check, creds, parts_by_id)
        print(f"[TIME] Lấy XML đính kèm: {time.perf_counter() - t0:.1f}s")
    if skipped_items:
        print(f"[SKIPPED SUMMARY] Email quet duoc nhung khong tai: {len(skipped_items)}")
        for skipped in skipped_items:
            eid = skipped.get("email_id", "UNKNOWN")
            print(
                "[SKIPPED] "
                f"EMAIL={eid} | "
                f"SUPPLIER={skipped.get('supplier')} | "
                f"REASON={skipped.get('reason')} | "
                f"SUBJECT={skipped.get('subject')}"
            )
            if eid != "UNKNOWN":
                print(f"[SKIPPED] OPEN=https://mail.google.com/mail/u/0/#all/{eid}")
    print(f"\n[STEP 3] Số lượng Link cần xử lý tải: {total_links}")

    sem = asyncio.Semaphore(MAX_CONCURRENT)
    success = 0
    failed = []
    t0 = time.perf_counter()

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)

        # Mỗi link cần 1 phiên trình duyệt riêng. Chỉ mở phiên khi tới lượt (tối đa MAX_CONCURRENT phiên
        # cùng lúc) - trước đây mở phiên cho TẤT CẢ link ngay từ đầu, hàng trăm link sẽ ngốn rất nhiều RAM.
        # Dùng semaphore riêng: `sem` được các hàm process_* tự giữ bên trong, giữ lồng 2 lần sẽ bị kẹt.
        ctx_sem = asyncio.Semaphore(MAX_CONCURRENT)

        async def wrapper(item):
            async with ctx_sem:
                await _process_link(item)

        async def _process_link(item):
            nonlocal success
            context = await browser.new_context(accept_downloads=True)
            eid = item.get("email_id", "UNKNOWN")
            ok = False
            err = ""
            saved_count = 0
            try:
                # TRUYỀN EMAIL_ID VÀ SUPPLIER VÀO TẤT CẢ CÁC HÀM
                if item["supplier"] == "easyinvoice":
                    saved_count = await process_easyinvoice(context, item["url"], sem, eid, item["supplier"])
                elif item["supplier"] == "meinvoice":
                    saved_count = await process_meinvoice(context, item["url"], sem, eid, item["supplier"])
                elif item["supplier"] == "pvoil":
                    saved_count = await process_pvoil(context, item["url"], sem, eid, item["supplier"])
                elif item["supplier"] == "smartsign":
                    saved_count = await process_smartsign(context, item["url"], sem, eid, item["supplier"])
                elif item["supplier"] == "fast":
                    saved_count = await process_fast(context, item["url"], sem, eid, item["supplier"])
                elif item["supplier"] == "ehoadon":
                    saved_count = await process_ehoadon(context, item["url"], sem, eid, item["supplier"])
                elif item["supplier"] == "vnpt":
                    saved_count = await process_vnpt(context, item["url"], sem, eid, item["supplier"])
                else:
                    saved_count = 0
                if not saved_count:
                    raise Exception("Không có file PDF nào được lưu.")
                success += saved_count
                ok = True
                if KEEP_XML and item["supplier"] in WEB_XML_SUPPLIERS:
                    await fetch_web_xml(context, item, sem)   # --keep-xml: lấy thêm XML (tự bắt lỗi)
            except Exception as e:
                err = repr(e)
                failed.append(item)
                if KEEP_XML and item["supplier"] in WEB_XML_SUPPLIERS:
                    await fetch_web_xml(context, item, sem)   # PDF lỗi nhưng vẫn lấy XML nếu được (tự bắt lỗi)
            finally:
                await context.close()

            files = pop_tracked_files(eid)
            if ok:
                record(eid, DONE, supplier=item["supplier"], saved_count=saved_count, files=files,
                       subject=item.get("subject"), mail_date=item.get("mail_date"))
            else:
                record(eid, FAILED, supplier=item["supplier"], reason=err[:300], files=files,
                       subject=item.get("subject"), mail_date=item.get("mail_date"))

        if link_items:
            SEQUENTIAL_SUPPLIERS = {"easyinvoice"}

            sequential = [i for i in link_items if i["supplier"] in SEQUENTIAL_SUPPLIERS]
            parallel   = [i for i in link_items if i["supplier"] not in SEQUENTIAL_SUPPLIERS]

            if parallel:
                await asyncio.gather(*[wrapper(i) for i in parallel])

            for item in sequential:
                await wrapper(item)

        await browser.close()

    print(f"\n[STEP 3 DONE] Đã tải thành công qua Link: {success}/{total_links}")
    print(f"[TIME] STEP 3 (tải qua link): {time.perf_counter() - t0:.1f}s")

    if failed:
        print("\n[STEP 4] Các email xử lý lỗi:")
        for f in failed:
            eid = f.get("email_id", "UNKNOWN")
            print(f"----- [{f.get('supplier').upper()}] -----")
            print(f"Tiêu đề : {f.get('subject')}")
            print(f"Email ID: {eid}")
            print(f"Link Web: {f.get('url')}")
            if eid != "UNKNOWN":
                print(f"Mở email: https://mail.google.com/mail/u/0/#all/{eid}")
            print("-------------------------\n")
    else:
        print("\n[STEP 4] Tất cả link hóa đơn đã được xử lý thành công!")

    # ---------- STEP 5: PETROLIMEX (xử lý sau cùng, khi các email khác đã xong) ----------
    print(f"\n[STEP 5] Xử lý email Petrolimex (chạy sau cùng): {len(petro_items)} email cần xử lý")
    petro_stats = process_petro_emails(gmail_service, state, petro_items, petro_label_ids, record)
    if petro_items:
        print(
            f"[STEP 5 DONE] Trích xuất đủ: {petro_stats['ok']} | Thiếu dữ liệu: {petro_stats['missing']} | "
            f"Không phải hóa đơn Petro: {petro_stats['not_petro']} | Lỗi đọc mail: {petro_stats['error']}"
        )
    else:
        print("[STEP 5 DONE] Không có email Petro mới.")
    # Excel dựng lại đầy đủ từ sổ (gồm cả các lần chạy trước); chỉ ghi lại khi có dữ liệu mới hoặc file bị mất
    dedupe = {"total": 0, "new": 0}
    if state.petro_count():
        dedupe = _petro_dedupe(state, verbose=True)   # lọc trùng: trong sổ + so với file tổng
    if petro_stats["saved"] or dedupe["new"] or (state.petro_count() and not os.path.exists(PETRO_OUTPUT)):
        _export_petro(state)

    # ---------- STEP 6: PETROLIMEX - TẢI PDF (cần nhập CAPTCHA thủ công, chạy cuối cùng) ----------
    if args.no_petro_download:
        print("\n[STEP 6] Bỏ qua tải PDF Petro (--no-petro-download).")
    else:
        print("\n[STEP 6] Tải PDF hóa đơn Petrolimex (chạy cuối cùng)...")
        dl = await run_petro_downloads(state, prompt=True)
        if dl["done"]:
            _export_petro(state)   # cập nhật cột "File PDF" trong Excel
    _petro_sync_master(state)      # ghi hóa đơn đã tải vào file tổng (kể cả khi bỏ qua bước tải, để bù các lần trước)

    # ---------- XUẤT EXCEL CÁC MỤC LỖI / BỎ QUA / CHƯA XỬ LÝ ----------
    xml_rows = xml_report_rows(meta_by_id) if KEEP_XML else []
    _export_issues(state, issues, link_url_by_id, xml_rows)

    if KEEP_XML:
        with _XML_LOCK:
            st_all = list(_XML_STATUS.values())
        n_ok = sum(1 for st in st_all if st["files"] and not st["errors"])
        n_err = sum(1 for st in st_all if st["errors"])
        n_none = sum(1 for st in st_all if not st["files"] and not st["errors"])
        print(f"[KEEP-XML] Lần này đã lưu {len(_XML_SAVED)} file XML vào: {_xml_dir()}")
        print(f"[KEEP-XML] Tổng kết XML: có XML {n_ok} email | không có XML {n_none} email | lỗi XML {n_err} email")
        for r in xml_rows:
            print(f"[KEEP-XML] - {r['status']} | EMAIL={r['email_id']} | LOAI={r['kind']}"
                  + (f" | PDF={r['pdf']}" if r["pdf"] else "") + (f" | LOI={r['error']}" if r["error"] else ""))

    # ---------- TỔNG KẾT LẦN CHẠY ----------
    counts = state.counts()
    print("\n[STATE] ===== TỔNG KẾT LẦN CHẠY =====")
    print(
        f"[STATE] Lần này: DONE {run_counts.get(DONE, 0)} | FAILED {run_counts.get(FAILED, 0)} | "
        f"SKIPPED {run_counts.get(SKIPPED, 0)} | Bỏ qua vì đã xử lý trước đó: {result.get('skipped_known', 0)}"
    )
    print("[STATE] Toàn sổ : " + " | ".join(f"{s}: {counts.get(s, 0)}" for s in ALL_STATUSES))
    if run_counts.get(FAILED, 0):
        print("[STATE] Email lỗi sẽ được tự thử lại ở lần chạy sau. Xem chi tiết: python main.py --report" + (f" --profile {args.profile}" if args.profile else ""))


async def main(args):
    ensure_dir(DOWNLOAD_DIR)
    creds = get_creds()
    gmail_service = build("gmail", "v1", credentials=creds)
    state = _open_state()
    try:
        await _pipeline(args, creds, gmail_service, state)
    finally:
        _export_csv(state)
        state.close()

async def main_download_petro(args):
    """Chế độ --download-petro: chỉ tải PDF Petro còn thiếu, không quét Gmail."""
    state = _open_state()
    try:
        stats = await run_petro_downloads(state, prompt=False)
        if stats["done"]:
            _export_petro(state)
            _export_csv(state)
        _petro_sync_master(state)
        _export_issues(state, {}, {})   # các hóa đơn Petro còn chưa tải được PDF
    finally:
        state.close()

class ScanError(Exception):
    """Lỗi khiến lần chạy không quét được Gmail (chương trình kết thúc với mã lỗi 1)."""


def cli():
    """Điểm vào của chương trình. Luôn in đường dẫn file log ở cuối, kể cả khi chạy lỗi."""
    try:
        _cli()
    except ScanError as e:
        print(f"[ERROR] {e}")
        print("[ERROR] Lần chạy này DỪNG, chưa tải được gì. Kiểm tra kết nối mạng / đăng nhập Gmail rồi chạy lại.")
        sys.exit(1)
    finally:
        if LOG_PATH:
            print(f"[LOG] Log lần chạy này: {LOG_PATH}")

def _cli():
    """Đọc tham số rồi chọn chế độ chạy."""
    args = parse_args()
    configure_runtime(args)
    if run_management_commands(args):
        return
    if args.download_petro:
        asyncio.run(main_download_petro(args))
        return
    asyncio.run(main(args))

if __name__ == "__main__":
    cli()
